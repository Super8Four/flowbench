"""Offline PDF and Word reports from a single, privacy-aware archive snapshot."""

import copy
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path
import re
import threading
from xml.sax.saxutils import escape
import xml.etree.ElementTree as ET

from .annotations import identifier
from .flent import normalize as normalize_flent, validate_flent
from .metrics import comparison, effective_settings, measurements, series
from .version import __version__
from .units import format_mbps, format_rate, rate_scale, run_color

LOG_LIMIT = 20000
MODEL_BYTES = 32 * 1024 * 1024
CHART_POINTS = 500000
REPORT_PAGES = 500
KNOWN_DIRECTIONS = {"upload", "download", "bidir"}
BRAND = Path(__file__).resolve().parent / "static" / "brand" / "flowbench-logo.png"
FONT_LOCK = threading.Lock()
SETTINGS = ("role", "target", "port", "bind", "protocol", "mode", "duration", "streams", "bandwidth", "ip_version")
LIMITATIONS = [
    "These captured tests describe measurements at their recorded times. Matching settings do not establish identical network conditions or prove causation.",
    "Throughput charts use elapsed seconds from each traffic start. Only measured intervals or points are drawn; missing spans are left blank.",
    "ICMP echo RTT includes probe scheduling and differs from application latency. Blocked, failed or unavailable probes are not zero latency.",
    "Endpoint health is local unless explicitly identified otherwise. A plain iPerf peer does not expose remote CPU or interface counters.",
    "Route replies and bounded MTU checks do not prove a throughput bottleneck or complete path MTU discovery.",
]


def _finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _clean(value):
    """Remove invalid control characters without interpreting user text as markup."""
    if isinstance(value, str):
        return "".join(c for c in value if c in "\n\t" or ord(c) >= 32 and not 0xD800 <= ord(c) <= 0xDFFF)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, dict):
        return {str(key): _clean(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_clean(item) for item in value]
    return value


def _private_diagnostics(value, key=""):
    """Default deny captured text: errors and tool output can contain endpoints."""
    safe_states = {"completed", "captured", "incomplete", "partial", "failed", "timeout", "unavailable", "disabled", "no_reply", "error",
                   "cancelled", "reply", "missing", "blocked", "permission_blocked", "skipped", "confirmed",
                   "unconfirmed", "separate", "detected local", "local", "remote", "manual", "linux", "windows"}
    safe_keys = set("schema_version latency endpoint_health path_context enabled metric baseline loaded samples summary mean_increase_ms sample_limit interval_seconds sample_limit_reached status timestamp elapsed_seconds rtt_ms rtt_upper_bound_ms note error errors exit_code local remote provenance before after delta cpu cpu_utilization_percent cpu_busy_percent interfaces counters total_ticks idle_ticks captured_at rx_bytes rx_errors rx_drops tx_bytes tx_errors tx_drops requested_ip_version probe_ip_version port protocol mtu route tested_packet_sizes largest_confirmed_packet_bytes packet_bytes hops hop responded observed observations concurrency overlaps incoming outgoing ids attempted replies missing numeric_samples min_ms max_ms mean_ms median_ms stddev_ms p95_ms phase started_at finished_at".split())
    if isinstance(value, dict):
        return {name: _private_diagnostics(item, name) for name, item in _without_endpoint_context(value).items() if name in safe_keys
                if name not in {"raw", "raw_output", "raw_stdout", "raw_stderr", "command", "executable", "address", "host", "target", "bind", "source", "interface", "name"}}
    if isinstance(value, list):
        return [_private_diagnostics(item, key) for item in value]
    if isinstance(value, str):
        if key == "phase" and value in {"before_traffic", "observed_after_start", "after_traffic"}:
            return value
        if key in {"timestamp", "captured_at", "started_at", "finished_at"}:
            try:
                parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
                if parsed.tzinfo is not None:
                    return parsed.isoformat()
            except ValueError:
                pass
        return value if key in {"status", "provenance", "platform"} and value in safe_states else "[redacted]"
    return _clean(value)


def _without_endpoint_context(value):
    """Remove sensitive endpoint subtrees before any generic report projection."""
    if isinstance(value, dict):
        if ({"sockets", "control_peers", "destinations"} <= value.keys()
                or value.get("source") == "wsgi_remote_addr" and "peer" in value):
            return {}
        return {key: _without_endpoint_context(item) for key, item in value.items()
                if key not in {"endpoints", "endpoint_view", "initiator_context", "web_request", "browser_peer"}}
    if isinstance(value, list):
        return [_without_endpoint_context(item) for item in value]
    return value


def _report_endpoints(record, redact):
    from .endpoints import endpoint_view, validate_context
    captured = endpoint_view(record)
    context = copy.deepcopy(record.get("initiator_context"))
    if context is not None:
        validate_context(context)
    if not redact:
        return _clean(captured), _clean(context)
    # Exact projection: even numeric identifiers and unknown future fields are
    # private. Raw opt-in never changes this boundary.
    sockets, destinations = captured.get("sockets", []), captured.get("destinations", [])
    points = [item.get(side) or {} for item in sockets for side in ("local", "remote")]
    points += [item.get("endpoint") or {} for item in destinations]
    summary = {"schema_version": 1, "status": captured["status"], "provenance": captured["provenance"],
               "identifiers_redacted": True, "socket_count": len(sockets),
               "control_peer_count": len(captured.get("control_peers", [])), "destination_count": len(destinations),
               "families": sorted({point["family"] for point in points if point.get("family") in (4, 6)}),
               "protocols": sorted({item["protocol"] for item in sockets if item.get("protocol") in {"tcp", "udp"}})}
    private_context = {"schema_version": 1, "status": "recorded" if context else "not_recorded",
                       "identifiers_redacted": True}
    return summary, private_context


def _private_preflight(value):
    """Keep numbers and known clock states; never expose provider-supplied text."""
    if not isinstance(value, dict):
        return None
    def state(value, choices):
        return value if isinstance(value, str) and value in choices else "[redacted]" if value else None
    def utc(value):
        if not isinstance(value, str):
            return None
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
            return parsed.astimezone(timezone.utc).isoformat() if parsed.tzinfo else None
        except ValueError:
            return None
    result = {"schema_version": 1, "policy": state(value.get("policy"), {"once_per_test"}),
              "status": state(value.get("status"), {"pending", "captured", "partial", "cancelled", "not_applicable"}),
              "phase": state(value.get("phase"), {"before_traffic", "observed_after_start"}),
              "started_at": utc(value.get("started_at")), "finished_at": utc(value.get("finished_at")),
              "elapsed_seconds": value.get("elapsed_seconds") if _finite(value.get("elapsed_seconds")) else None}
    quality = value.get("time_quality")
    if isinstance(quality, dict):
        result["time_quality"] = {
            "schema_version": 1,
            "status": state(quality.get("status"), {"pending", "captured", "status_only", "malformed", "unavailable", "unsynchronized", "timeout", "error", "cancelled", "permission_blocked"}),
            "observed_at": utc(quality.get("observed_at")), "clock_source": "system_utc", "ntp_verified": False,
            "provider": state(quality.get("provider"), {"chrony", "windows_time", "systemd"}),
            "source": "[redacted]" if quality.get("source") else None,
            "synchronized": quality.get("synchronized") if isinstance(quality.get("synchronized"), bool) else None,
            "uncertainty_kind": state(quality.get("uncertainty_kind"), {"root_dispersion_proxy"}),
            "last_sync_at": utc(quality.get("last_sync_at")), "note": "[redacted]" if quality.get("note") else None}
        for key in ("offset_seconds", "uncertainty_seconds", "root_delay_seconds", "root_dispersion_seconds", "query_elapsed_seconds"):
            result["time_quality"][key] = quality.get(key) if _finite(quality.get(key)) else None
    return result


def _private_route(value):
    """Route geometry/location is sensitive, including unknown future fields."""
    if not isinstance(value, dict):
        return None
    result = {key: value[key] for key in ("status", "note", "phase", "started_at", "finished_at", "elapsed_seconds", "exit_code") if key in value}
    result["hops"] = []
    for hop in value.get("hops", []):
        kept = {key: hop[key] for key in ("hop", "responded", "observed", "rtt_ms", "rtt_upper_bound_ms") if key in hop}
        if isinstance(hop.get("observations"), list):
            kept["observations"] = [{key: item[key] for key in ("rtt_ms", "rtt_upper_bound_ms") if key in item}
                                    for item in hop["observations"] if isinstance(item, dict)]
        result["hops"].append(kept)
    return _private_diagnostics(result)


def _bounded_raw(record):
    result = {}
    for key in ("raw_stdout", "raw_stderr"):
        text = _clean(record.get(key) or "")
        result[key] = text[:LOG_LIMIT]
        result[key + "_truncated"] = len(text) > LOG_LIMIT
    # Native tool evidence can include sensitive context, so it shares the raw opt-in.
    native = record.get("native_result") or record.get("flent") or record.get("result")
    if native:
        text = json.dumps(_clean(native), indent=2, ensure_ascii=False)
        result["native_result"] = text[:LOG_LIMIT]
        result["native_result_truncated"] = len(text) > LOG_LIMIT
    return result


def _private_check(value, key=""):
    """Default-deny standalone check context, including unknown numeric fields."""
    fields = set("schema_version status started_at finished_at elapsed_seconds time_quality adapter resolver resolution attempts summary error note mode port address_family transport query_index sequence duration_ms resolver_port resolver_family record_type resolver_mode count timeout_ms rcode flags question answers authority additional name type class ttl rdata qtype timing query_elapsed_seconds seconds rtt_ms rtt_upper_bound_ms transmitted echo_replies requested_queries attempted_queries transport_attempts responses outcomes requested attempted unattempted confirmed_transmitted confirmed_nonresponses indeterminate_attempts nonresponse_percent rtt numeric_samples min_ms max_ms mean_ms median_ms stdev_ms stddev_ms p95_ms records_truncated output_truncated".split())
    states = set("pending running completed captured partial failed interrupted cancelled unavailable blocked permission_blocked error timeout malformed reply no_reply unreachable missing answer nodata nxdomain dns_error truncated refused network_error resolution_error udp tcp system custom ipv4 ipv6 auto 4 6 linux win32 windows ICMP DNS dnspython native_ping A AAAA CNAME MX NS PTR SOA TXT SRV CAA HTTPS SVCB IN NOERROR NXDOMAIN SERVFAIL REFUSED FORMERR NOTIMP YXDOMAIN YXRRSET NXRRSET NOTAUTH NOTZONE QR AA TC RD RA AD CD".split())
    if key == "time_quality":
        return (_private_preflight({"time_quality": value}) or {}).get("time_quality")
    if key == "adapter" and isinstance(value, dict):
        return {"name": value.get("name") if value.get("name") in {"dnspython", "native_ping"} else "[redacted]",
                "version": "[redacted]" if value.get("version") else None,
                "platform": value.get("platform") if value.get("platform") in {"linux", "win32", "darwin"} else "[redacted]"}
    if isinstance(value, dict):
        return {name: _private_check(item, name) for name, item in _without_endpoint_context(value).items() if name in fields or key == "outcomes" and name in states}
    if isinstance(value, list):
        return [_private_check(item, key) for item in value]
    if isinstance(value, str):
        if key in {"started_at", "finished_at"}:
            try:
                parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
                return parsed.astimezone(timezone.utc).isoformat() if parsed.tzinfo else None
            except ValueError:
                return None
        return value if value in states and key in {"status", "mode", "resolver_mode", "address_family", "transport", "rcode", "flags", "record_type", "qtype", "type", "class"} else "[redacted]"
    return _clean(value)


def _check_report(manager, records, projects, redact, include_raw, expected_fingerprint):
    """Build standalone diagnostic evidence without iPerf settings or metrics."""
    if any(record.get("kind") != "diagnostic-check" for record in records):
        raise ValueError("Choose diagnostic checks separately from throughput tests.")
    if len({record["check_type"] for record in records}) != 1:
        raise ValueError("Compare DNS checks and ICMP checks separately.")
    model = {"schema_version": 1, "kind": "flowbench-report", "report_type": "diagnostic-check",
             "generated_at": datetime.now(timezone.utc).isoformat(), "version": __version__,
             "title": "Network diagnostic report", "redacted": redact, "include_raw": include_raw,
             "runs": [], "comparisons": [], "common_window": None,
             "limitations": ["This report records standalone diagnostic observations. Failed, partial, cancelled and interrupted captures retain their actual outcomes; unavailable measurements are not zero.",
                             "DNS responses describe the selected resolver and transport at capture time. Negative or error responses are observations, not proof that every resolver or network path behaves the same way.",
                             "ICMP echo nonresponse is not UDP or application packet loss. Filtering, rate limits, local permissions and unknown sends can prevent a reply. Bounds such as <1 ms are not numeric zero RTT.",
                             "UTC timestamps use the host system clock and are not independently NTP verified. Local provider synchronization and root dispersion are not guaranteed clock accuracy.",
                             "No throughput series or automatic route capture is inferred from these checks.",
                             "PDF embeds offline Bitstream Vera fonts; unsupported characters become a question mark. Word retains Unicode text using its font fallback.",
                             "Redaction removes names, DNS answers, resolver and endpoint addresses, project text, notes, raw output and time-source context." if redact else "Endpoint, resolver, DNS answer and user context may be sensitive. Review before sharing."]}
    def drop_raw(value):
        if isinstance(value, dict):
            return {k: drop_raw(v) for k, v in value.items() if k not in {"raw", "raw_output", "raw_stdout", "raw_stderr", "command", "executable"}}
        return [drop_raw(v) for v in value] if isinstance(value, list) else value
    for index, record in enumerate(records, 1):
        settings = copy.deepcopy(record["settings"])
        project = projects.get(settings.get("project_id"))
        result = copy.deepcopy(record["result"])
        endpoints, initiator = _report_endpoints(record, redact)
        row = {"id": record["id"], "kind": "diagnostic-check", "check_type": record["check_type"],
               "display_index": index, "color": run_color(record), "label": "Check " + str(index) if redact else record.get("label") or settings.get("label") or "Check " + str(index),
               "created_at": record["created_at"], "finished_at": record["finished_at"], "status": record["status"],
               "settings": _private_check(settings) if redact else settings,
               "result": _private_check(result) if redact else drop_raw(result), "metrics": {}, "series": [],
               "endpoints": endpoints, "initiator_context": initiator,
               "project": "[redacted]" if redact and project else copy.deepcopy(project),
               "tags": {} if redact else copy.deepcopy(settings.get("tags", {})),
               "notes": "[redacted]" if redact and settings.get("notes") else settings.get("notes", ""),
               "annotations": [{"timestamp": n.get("timestamp"), "seconds": n.get("seconds"), "text": "[redacted]"} for n in record.get("annotations", [])] if redact else copy.deepcopy(record.get("annotations", [])),
               "provenance": {"engine": "DNS query" if record["check_type"] == "dns" else "ICMP echo",
                              "imported": bool(record.get("imported")), "output_format": "typed diagnostic evidence",
                              "lineage": copy.deepcopy(record.get("lineage"))}}
        if include_raw:
            row["raw_evidence"] = {"status": "removed by redaction"} if redact else _bounded_raw(record)
        model["runs"].append(_clean(row))
    # The shared check comparison adapter is supplied by the typed check module.
    if len(records) > 1:
        from .metrics import check_comparison as compare_checks
        for after in records[1:]:
            captured = compare_checks(records[0], after)
            metrics = [{"metric": item["label"] + " (" + item["unit"] + ")", "unit": item["unit"], "before": item["before"], "after": item["after"], "delta": item["delta"]} for item in captured["rows"]]
            if captured.get("outcomes"):
                for state in sorted(captured["outcomes"]["before"]):
                    metrics.append({"metric": "DNS " + state + " (count)", "before": captured["outcomes"]["before"][state],
                                    "after": captured["outcomes"]["after"][state], "unit": "count", "delta": captured["outcomes"]["after"][state] - captured["outcomes"]["before"][state]})
            details = [("Baseline capture", captured["before"]["status"] + " / " + captured["before"]["result_status"]),
                       ("Selected capture", captured["after"]["status"] + " / " + captured["after"]["result_status"])]
            if not redact and captured.get("answer_sets"):
                details.extend(("DNS answer sets " + key, _summary(value)) for key, value in captured["answer_sets"].items())
            model["comparisons"].append({"before_id": captured["before_id"], "after_id": captured["after_id"],
                                        "caveat": " ".join(captured["caveats"]), "settings_differences": captured["settings_differences"],
                                        "metrics": metrics, "details": details})
    content = {key: value for key, value in model.items() if key != "generated_at"}
    serialized = json.dumps(content, ensure_ascii=False, sort_keys=True, allow_nan=False).encode()
    if len(serialized) > MODEL_BYTES:
        raise ValueError("Report exceeds the 32 MiB normalized-data resource budget. Split the selection into smaller reports.")
    model["fingerprint"] = hashlib.sha256(serialized).hexdigest()
    if expected_fingerprint is not None and expected_fingerprint != model["fingerprint"]:
        raise ValueError("Saved evidence changed; preview the report again before downloading.")
    return model


def prepare_report(manager, data):
    """Read completed leaf tests within explicit resource budgets; never mutate evidence."""
    if not isinstance(data, dict) or set(data) - {"ids", "redact", "include_raw", "fingerprint"}:
        raise ValueError("Send ids, redact and include_raw as report options.")
    ids = data.get("ids")
    if not isinstance(ids, list) or not ids or len(set(str(i) for i in ids)) != len(ids):
        raise ValueError("Choose distinct completed tests.")
    ids = [identifier(item, "Report test ID") for item in ids]
    redact, include_raw = data.get("redact", True), data.get("include_raw", False)
    if not isinstance(redact, bool) or not isinstance(include_raw, bool):
        raise ValueError("Redact and include_raw must be true or false.")
    expected_fingerprint = data.get("fingerprint")
    if expected_fingerprint is not None and (not isinstance(expected_fingerprint, str) or not re.fullmatch(r"[a-f0-9]{64}", expected_fingerprint)):
        raise ValueError("Invalid report preview fingerprint.")
    with manager.lock:
        records, point_count = [], 0
        for run_id in ids:
            record = manager.read(run_id)
            if "lineage" in record:
                from .reruns import validate_lineage
                validate_lineage(record["lineage"])
            from .endpoints import validate_endpoints, validate_context
            if "endpoints" in record:
                validate_endpoints(record["endpoints"], record)
            if "initiator_context" in record:
                validate_context(record["initiator_context"])
            if record.get("children") or record.get("kind") in {"protocol-job", "matrix-job"}:
                raise ValueError("Choose individual tests within a protocol or matrix batch.")
            if record.get("kind") == "diagnostic-check":
                from .checks import validate_check_record
                validate_check_record(record)
                if record.get("status") not in {"completed", "partial", "failed", "cancelled", "interrupted"} or not record.get("finished_at"):
                    raise ValueError("Reports require terminal diagnostic checks.")
                records.append(manager.export(run_id) if include_raw and not redact else record)
                continue
            if record.get("status") != "completed" or not record.get("finished_at"):
                raise ValueError("Reports require completed tests.")
            if "preflight" in record:
                from .preflight import validate_preflight
                validate_preflight(record["preflight"])
            route = record.get("route")
            if isinstance(route, dict) and ("geolocation" in route or any("locations" in hop for hop in route.get("hops", []) if isinstance(hop, dict))):
                from .route import validate_route
                validate_route(route)
            if record.get("kind") == "flent-run":
                # Legacy archive rows can claim completion despite missing load
                # flows. Revalidate source evidence without changing that row.
                evidence = record.get("flent")
                try:
                    validate_flent(evidence)
                    parsed = normalize_flent(evidence.get("data"), evidence.get("version"))
                    complete = evidence.get("status") == "captured" and parsed["quality"]["complete"]
                except ValueError:
                    complete = False
                if not complete:
                    raise ValueError("Reports require a complete native Flent RRUL capture with all eight TCP load flows and four RTT series.")
            if not measurements(record):
                raise ValueError("Reports require captured performance measurements; this workload has no reportable metrics.")
            point_count += len(series(record)) + sum(len(item.get("points", [])) for item in (record.get("flent") or {}).get("series", []))
            if point_count > CHART_POINTS:
                raise ValueError("Report exceeds the 500000 measured-point resource budget. Split the selection into smaller reports.")
            records.append(manager.export(run_id) if include_raw and not redact else record)
        catalog = manager.storage.metadata("projects") or {"projects": []}
        projects = {item["id"]: item for item in catalog.get("projects", [])}
    if any(record.get("kind") == "diagnostic-check" for record in records):
        return _check_report(manager, records, projects, redact, include_raw, expected_fingerprint)
    model = {"schema_version": 1, "kind": "flowbench-report", "generated_at": datetime.now(timezone.utc).isoformat(),
             "version": __version__, "title": "Network performance report", "redacted": redact,
             "include_raw": include_raw, "runs": [], "comparisons": [], "limitations": list(LIMITATIONS)}
    if redact:
        model["limitations"].append("Redaction removes endpoint addresses and hostnames, project text, labels, tags, user notes and raw evidence. Review the preview before sharing.")
    else:
        model["limitations"].append("This report contains endpoint and user context. Review its contents before sharing outside the intended audience.")
    model["limitations"].append("PDF uses embedded Bitstream Vera fonts; characters outside their supported glyph set are replaced with a question mark. Word retains Unicode text.")
    metric_names = {}
    private_directions = {}
    def report_direction(value):
        if not redact or value in KNOWN_DIRECTIONS:
            return value
        if value not in private_directions:
            private_directions[value] = "redacted_direction_" + str(len(private_directions) + 1)
        return private_directions[value]
    def report_series(record):
        captured = series(record)
        if not redact:
            return captured
        # Imported stream/direction labels are arbitrary user text, not enums.
        return [dict({key: point[key] for key in ("seconds", "mbps", "start") if key in point},
                     direction=report_direction(point.get("direction"))) for point in captured]
    def report_window(window):
        if not redact:
            return window
        result = copy.deepcopy(window)
        for run in result.get("runs", []):
            for direction in run.get("directions", []):
                direction["direction"] = report_direction(direction.get("direction"))
        return result
    for key in sorted({key for record in records for key in measurements(record)}):
        if not redact or re.fullmatch(r"(?:upload|download)_(?:mbps|jitter_ms|lost_percent|lost_packets|packets|retransmits)", key):
            metric_names[key] = key
        else:
            suffix = "_mbps" if key.endswith("_mbps") else "_rtt_ms" if key.endswith("_rtt_ms") else ""
            metric_names[key] = "external_metric_" + str(len(metric_names) + 1) + suffix
    for index, record in enumerate(records, 1):
        settings = effective_settings(record)
        endpoints, initiator = _report_endpoints(record, redact)
        project = projects.get(settings.get("project_id"))
        diagnostics, route = copy.deepcopy(record.get("diagnostics")), copy.deepcopy(record.get("route"))
        preflight = copy.deepcopy(record.get("preflight"))
        if not include_raw and not redact:
            def drop_raw(value):
                if isinstance(value, dict):
                    return {k: drop_raw(v) for k, v in value.items() if k not in {"raw", "raw_output", "raw_stdout", "raw_stderr", "command", "executable", "detail"}}
                return [drop_raw(v) for v in value] if isinstance(value, list) else value
            diagnostics, route = drop_raw(diagnostics), drop_raw(route)
            preflight = drop_raw(preflight)
        row = {"id": record["id"], "display_index": index, "color": run_color(record), "label": "Test " + str(index) if redact else record.get("label") or settings.get("label") or "Test " + str(index),
               "created_at": record["created_at"], "finished_at": record["finished_at"], "status": record["status"],
               "settings": {k: settings.get(k) for k in SETTINGS}, "metrics": {metric_names[k]: v for k, v in measurements(record).items()}, "series": report_series(record),
               "project": "[redacted]" if redact and project else copy.deepcopy(project),
               "tags": {} if redact else copy.deepcopy(record.get("tags", settings.get("tags", {}))),
               "notes": "[redacted]" if redact and (record.get("notes") or settings.get("notes")) else record.get("notes", settings.get("notes", "")),
               "annotations": [{"timestamp": note.get("timestamp"), "seconds": note.get("seconds"), "text": "[redacted]"}
                               for note in record.get("annotations", [])] if redact else copy.deepcopy(record.get("annotations", [])),
               "diagnostics": _private_diagnostics(diagnostics) if redact else diagnostics,
               "route": _private_route(route) if redact else route,
               "preflight": _private_preflight(preflight) if redact else preflight,
               "endpoints": endpoints, "initiator_context": initiator,
               "provenance": {"engine": record.get("engine", "iperf3"), "engine_version": record.get("iperf_version"),
                              "output_format": record.get("output_format"), "imported": bool(record.get("imported")),
                              "lineage": copy.deepcopy(record.get("lineage")),
                              "concurrency": _private_diagnostics(record.get("concurrency")) if redact else copy.deepcopy(record.get("concurrency"))}}
        if record.get("kind") == "flent-run":
            flent = record.get("flent") or {}
            row["workload"] = "Actual Flent RRUL with four upload and four download TCP flows"
            row["provenance"]["engine"] = "Flent"
            row["provenance"]["engine_version"] = flent.get("version")
            row["settings"]["streams"] = "RRUL four per direction"
            row["flent_series"] = [{"name": item.get("name"), "unit": item.get("unit") if item.get("unit") in {"Mbits/s", "ms"} else None,
                                    "points": [{"seconds": p.get("seconds"), "value": p.get("value")} for p in item.get("points", [])]}
                                   for item in flent.get("series", [])]
            row["flent_summaries"] = [{k: item.get(k) for k in ("name", "unit", "count", "missing", "mean", "min", "max")}
                                      for item in flent.get("summaries", [])]
            if redact:
                for number, item in enumerate(row["flent_series"], 1):
                    suffix = "_mbps" if item["unit"] == "Mbits/s" else "_rtt_ms" if item["unit"] == "ms" else ""
                    item["name"] = metric_label(metric_names.get("flent_" + str(item["name"]) + suffix, f"Series {number}"))
                for number, item in enumerate(row["flent_summaries"], 1):
                    suffix = "_mbps" if item["unit"] == "Mbits/s" else "_rtt_ms" if item["unit"] == "ms" else ""
                    item["name"] = metric_label(metric_names.get("flent_" + str(item["name"]) + suffix, f"Series {number}"))
                    item["unit"] = item["unit"] if item["unit"] in {"Mbits/s", "ms"} else None
        if redact:
            for key in ("target", "bind"):
                row["settings"][key] = "[redacted]" if settings.get(key) else settings.get(key)
            # Engine versions are captured free text and may come from imports.
            row["provenance"]["engine_version"] = "[redacted]" if row["provenance"].get("engine_version") else None
            row["provenance"]["engine"] = "Flent" if record.get("kind") == "flent-run" else "iperf3" if record.get("engine", "iperf3") == "iperf3" else "external workload"
            row["provenance"]["output_format"] = record.get("output_format") if record.get("output_format") in {"json", "text"} else "unknown"
        if include_raw:
            row["raw_evidence"] = {"status": "removed by redaction"} if redact else _bounded_raw(record)
        model["runs"].append(_clean(row))
    from .comparisons import common_window
    model["common_window"] = report_window(common_window(records))
    model["display"] = {"rate_axis": rate_scale([p["mbps"] * 1_000_000 for run in model["runs"] for p in run["series"] if _finite(p.get("mbps"))]
                                               + [p["value"] * 1_000_000 for run in model["runs"] for item in run.get("flent_series", []) if item["unit"] == "Mbits/s" for p in item["points"] if _finite(p.get("value"))])}
    model["limitations"].append("Summary metrics cover each complete captured run. A common elapsed window does not recompute these whole-run averages. Shorter traces stop at their actual last measurement; they are not extended to match longer tests.")
    model["limitations"].append("Saved colors are used consistently. Chart panels and legends paginate; colors can repeat and labels identify traces. Resource limits are 32 MiB of normalized report data, 500000 measured chart points, 500 PDF pages and 32 MiB of Word document content.")
    for after in records[1:]:
        report = comparison(records[0], after)
        if "common_window" in report:
            report["common_window"] = report_window(report["common_window"])
        # Packet counts indicate test volume rather than a universal quality gain.
        for item in report["metrics"]:
            item["metric"] = metric_names[item["metric"]]
            if item["metric"].endswith("_packets") and not item["metric"].endswith("lost_packets"):
                item["direction"] = "volume change" if item["delta"] else "unchanged"
        report["caveat"] = ("Settings differ. Measured changes do not establish an improvement or regression." if report["settings_differences"] else report["interpretation"])
        model["comparisons"].append(_clean(report))
    content = {key: value for key, value in model.items() if key != "generated_at"}
    serialized = json.dumps(content, ensure_ascii=False, sort_keys=True, allow_nan=False).encode()
    if len(serialized) > MODEL_BYTES:
        raise ValueError("Report exceeds the 32 MiB normalized-data resource budget. Split the selection into smaller reports.")
    model["fingerprint"] = hashlib.sha256(serialized).hexdigest()
    if expected_fingerprint is not None and expected_fingerprint != model["fingerprint"]:
        raise ValueError("Saved evidence changed; preview the report again before downloading.")
    return model


def display(value):
    if value is None:
        return "Unavailable"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, float):
        return f"{value:,.3f}".rstrip("0").rstrip(".")
    return str(value)


def metric_label(key):
    return key.replace("_", " ").replace("mbps", "rate").replace("jitter ms", "jitter ms").replace("lost percent", "loss %")


def metric_value(key, value):
    if key.endswith("_mbps") and _finite(value):
        return format_mbps(value)
    if key.endswith("(ms)") and _finite(value):
        return format(value, ".9g")
    return display(value)


def _clock_seconds(value, signed=False):
    if not _finite(value):
        return "Unavailable"
    return format(value, "+.9g" if signed and value else ".9g")


def _bandwidth(value):
    match = re.fullmatch(r"(\d+(?:\.\d+)?)([KMG]?)", str(value), re.I)
    if match:
        number = float(match[1]) * {"": 1, "K": 1000, "M": 1e6, "G": 1e9}[match[2].upper()]
        return "Unlimited (requested 0)" if number == 0 else format_rate(number)
    return display(value)


def _summary(value):
    text = display(value)
    return text if len(text) <= 800 else text[:800] + " [Report summary truncated; complete evidence is in the saved JSON or XML export]"


def _bounds(runs, native_unit=None):
    points = []
    for run in runs:
        if native_unit:
            points.extend({"seconds": point.get("seconds"), "value": point.get("value")} for item in run.get("flent_series", [])
                          if item.get("unit") == native_unit for point in item.get("points", []))
        else:
            points.extend({"seconds": point.get("seconds"), "value": point.get("mbps")} for point in run["series"])
    points = [p for p in points if _finite(p["seconds"]) and _finite(p["value"]) and p["seconds"] >= 0 and p["value"] >= 0]
    multiplier = 1 if native_unit == "ms" else 1_000_000
    scale = {"unit": "ms", "divisor": 1} if native_unit == "ms" else rate_scale([p["value"] * multiplier for p in points])
    return {"xmax": max(1, max((p["seconds"] for p in points), default=0)),
            "ymax": max(1, max((p["value"] * multiplier / scale["divisor"] for p in points), default=0)) * 1.1,
            "scale": scale, "multiplier": multiplier}


def _chart(runs, bounds=None):
    """Raster chart of individual measured spans, with no bridging of missing data."""
    from PIL import Image, ImageDraw, ImageFont
    import reportlab
    fontpath = Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"
    font = ImageFont.truetype(str(fontpath), 23)
    image = Image.new("RGB", (1500, 620), "white")
    draw = ImageDraw.Draw(image)
    all_points = [(index, point) for index, run in enumerate(runs) for point in run["series"]
                  if _finite(point.get("seconds")) and _finite(point.get("mbps")) and point["seconds"] >= 0 and point["mbps"] >= 0]
    if not all_points:
        draw.text((75, 230), "No measured interval series is available", fill="#1b1f22", font=font)
    else:
        bounds = bounds or _bounds(runs)
        xmax, ymax, scale = bounds["xmax"], bounds["ymax"], bounds["scale"]
        x0, y0, x1, y1 = 130, 75, 1435, 465
        for i in range(5):
            x, y = x0 + (x1 - x0) * i / 4, y1 - (y1 - y0) * i / 4
            draw.line((x0, y, x1, y), fill="#d9d9d9", width=1)
            draw.text((12, y - 12), f"{ymax * i / 4:.1f}", fill="#41494d", font=font)
            draw.text((x - 20, y1 + 15), f"{xmax * i / 4:.1f}", fill="#41494d", font=font)
        draw.text((12, 20), scale["unit"], fill="#1b1f22", font=font)
        draw.text((590, 520), "Elapsed seconds", fill="#1b1f22", font=font)
        draw.line((x0, y0, x0, y1, x1, y1), fill="#41494d", width=2)
        for index, point in all_points:
            color = runs[index]["color"]
            x = x0 + point["seconds"] / xmax * (x1 - x0)
            y = y1 - (point["mbps"] * 1_000_000 / scale["divisor"]) / ymax * (y1 - y0)
            start = point.get("start")
            if _finite(start) and 0 <= start < point["seconds"]:
                sx = x0 + start / xmax * (x1 - x0)
                if point.get("direction") == "download":
                    for pos in range(int(sx), int(x), 15):
                        draw.line((pos, y, min(pos + 8, x), y), fill=color, width=3)
                elif point.get("direction") in {"upload", "bidir"}:
                    draw.line((sx, y, x, y), fill=color, width=3)
            draw.ellipse((x - 3, y - 3, x + 3, y + 3), fill=color)
        draw.text((130, 575), "Trace labels and saved colors are listed below", fill="#41494d", font=font)
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def _check_sections(run):
    result, settings = run["result"], run["settings"]
    def value(item):
        return format(item, ".9g") if _finite(item) else _summary(item)
    kind = "DNS query" if run["check_type"] == "dns" else "ICMP echo nonresponse"
    context = [("Check type", kind), ("Saved capture status", run["status"]),
               ("Diagnostic outcome", result.get("status")),
               ("Saved check ID", run["id"]), ("Recorded UTC", run["created_at"] + " to " + run["finished_at"]),
               ("Capture UTC", display(result.get("started_at")) + " to " + display(result.get("finished_at"))),
               ("Elapsed seconds", value(result.get("elapsed_seconds"))),
               ("Adapter", _summary(result.get("adapter"))), ("Imported evidence", run["provenance"]["imported"])]
    if run["provenance"].get("lineage"):
        context.append(("Execution source lineage", _summary(run["provenance"]["lineage"])))
    sections = [("Check settings and provenance", ["Setting", "Recorded value"], context + [(key.replace("_", " "), value(item)) for key, item in settings.items() if key not in {"label", "notes", "tags", "project_id", "color"}]),
                ("Captured diagnostic summary", ["Metric", "Value"], [(key.replace("_", " "), value(item)) for key, item in result.get("summary", {}).items()])]
    quality = result.get("time_quality") or {}
    sections.append(("Clock quality", ["Evidence", "Recorded value"],
                     [("Verification", "Host system UTC; not independently NTP verified. Local synchronization is a provider report, not guaranteed clock accuracy."),
                      ("Status and observation UTC", display(quality.get("status")) + " | " + display(quality.get("observed_at"))),
                      ("Provider and source", display(quality.get("provider")) + " | " + display(quality.get("source"))),
                      ("Provider synchronized", quality.get("synchronized")),
                      ("Reported correction seconds", _clock_seconds(quality.get("offset_seconds"), signed=True)),
                      ("Root dispersion proxy seconds", _clock_seconds(quality.get("uncertainty_seconds")) + "; not guaranteed accuracy"),
                      ("Last synchronization UTC", quality.get("last_sync_at"))]))
    if run["check_type"] == "dns":
        sections.append(("Resolver observation", ["Evidence", "Recorded value"], [(key.replace("_", " "), value(item)) for key, item in (result.get("resolver") or {}).items()]))
        for ordinal, attempt in enumerate(result.get("attempts", []), 1):
            keys = ("sequence", "query_index", "transport", "status", "rcode", "flags", "question", "resolver_address", "resolver_port", "resolver_family", "started_at", "finished_at", "duration_ms", "records_truncated", "output_truncated", "error")
            sections.append((f"DNS transport attempt {ordinal}", ["Evidence", "Recorded value"],
                             [(key.replace("_", " "), value(attempt[key])) for key in keys if key in attempt]))
            records = []
            for group in ("answers", "authority", "additional"):
                for item in attempt.get(group, []):
                    records.append((group, item.get("name"), display(item.get("type")) + " / " + display(item.get("class")), item.get("ttl"), _summary(item.get("rdata"))))
            if records:
                sections.append((f"DNS response records {ordinal}", ["Section", "Name", "Type / class", "TTL", "Data"], records))
    else:
        sections.append(("Selected endpoint resolution", ["Evidence", "Recorded value"], [(key.replace("_", " "), value(item)) for key, item in (result.get("resolution") or {}).items()]))
        probes = []
        for ordinal, attempt in enumerate(result.get("attempts", []), 1):
            rtt = (_clock_seconds(attempt.get("rtt_ms")) if _finite(attempt.get("rtt_ms")) else
                   "<" + _clock_seconds(attempt["rtt_upper_bound_ms"]) if _finite(attempt.get("rtt_upper_bound_ms")) else "Unavailable")
            probes.append((ordinal, attempt.get("status"), rtt, attempt.get("transmitted"), attempt.get("echo_replies")))
        sections.append(("ICMP echo observations", ["Attempt", "Status", "RTT ms", "Transmitted", "Echo replies"], probes))
    if result.get("error") or result.get("note"):
        sections.append(("Capture notes", ["Evidence", "Recorded value"], [(key, _summary(result.get(key))) for key in ("error", "note") if result.get(key)]))
    return sections


def _endpoint_sections(run):
    evidence = run.get("endpoints") or {}
    def point(value):
        value = value or {}
        address = value.get("address")
        if not address:
            return "Unavailable"
        address = "[" + address + "]" if ":" in address else address
        return address + (":" + str(value["port"]) if value.get("port") is not None else " (port unavailable)")
    rows = [("Capture status / provenance", display(evidence.get("status")) + " / " + display(evidence.get("provenance"))),
            ("Interpretation", "Local and remote remain in each reporter's perspective, including reverse and bidirectional tests. Process/container namespace addresses do not establish a physical host, NAT egress or public IP.")]
    if evidence.get("identifiers_redacted"):
        rows.extend([("Endpoint identifiers", "Removed by redaction, including hostnames, socket IDs and ports"),
                     ("Data sockets / control peers / destinations", f"{evidence.get('socket_count', 0)} / {evidence.get('control_peer_count', 0)} / {evidence.get('destination_count', 0)}"),
                     ("Observed families / protocols", _summary(evidence.get("families", [])) + " / " + _summary(evidence.get("protocols", [])))])
    else:
        rows.append(("Execution hostname", evidence.get("execution_hostname") or "Not recorded; no current-host substitution"))
        for index, item in enumerate(evidence.get("sockets", []), 1):
            label = f"Data socket {index}: {item['reporter']} / {item['observer']} / {item.get('protocol') or 'protocol unavailable'}"
            rows.append((label, "Local " + point(item.get("local")) + "\nRemote " + point(item.get("remote")) + "\nSource: " + item["source"]))
        if not evidence.get("sockets"):
            rows.append(("Measured data sockets", "Not captured or unavailable; requested settings are not observed sockets"))
        for index, item in enumerate(evidence.get("control_peers", []), 1):
            rows.append((f"Control peer {index}: {item['reporter']}", point(item.get("peer")) + "; control connection, not a data-stream port"))
        for index, item in enumerate(evidence.get("destinations", []), 1):
            rows.append((f"Destination {index}: {item['purpose']}", point(item.get("endpoint")) + "\nReported host: " + display(item.get("reported_host")) + "\nSource: " + item["source"] + "; " + display(item.get("transport"))))
    context = run.get("initiator_context")
    if isinstance(context, dict) and context.get("identifiers_redacted"):
        request = "Identifiers removed by redaction" if context["status"] == "recorded" else "Not recorded"
    elif context:
        request = point(context.get("peer")) + " | Observed UTC " + context["observed_at"] + " | WSGI request peer"
    else:
        request = "Not recorded; no current-viewer substitution"
    rows.append(("Initiating HTTP peer (not traffic source)", request + ". Request context can identify a proxy peer; it does not identify the traffic generator or public client IP."))
    return [("Captured test endpoints", ["Evidence", "Recorded value"], rows)]


def _sections(run):
    """One shared narrative/table definition prevents PDF and Word drift."""
    if run.get("kind") == "diagnostic-check":
        sections = _check_sections(run)
        return sections[:1] + _endpoint_sections(run) + sections[1:]
    s = run["settings"]
    settings = [("Role and endpoint", f"{s.get('role')} | {display(s.get('target'))}:{display(s.get('port'))}\nBind: {s.get('bind') or 'Automatic'}"),
                ("Protocol and direction", f"{s.get('protocol')} | {s.get('mode')}"),
                ("Requested duration and streams", f"{display(s.get('duration'))} seconds | {display(s.get('streams'))} streams"),
                ("Requested bandwidth and IP version", f"{_bandwidth(s.get('bandwidth'))} | {display(s.get('ip_version'))}"),
                ("Recorded timestamps", "Started " + run["created_at"] + "\nFinished " + run["finished_at"]),
                ("Clock verification", "Host system UTC; not independently NTP verified. Provider synchronization is local status only."),
                ("Saved test ID", run["id"]),
                ("Engine and evidence", f"{display(run['provenance'].get('engine'))} | {display(run['provenance'].get('engine_version'))}\nOutput {display(run['provenance'].get('output_format'))} | Imported {display(run['provenance'].get('imported'))}")]
    if run["provenance"].get("concurrency"):
        settings.append(("Overlapping traffic", _summary(run["provenance"]["concurrency"])))
    if run["provenance"].get("lineage"):
        settings.append(("Execution source lineage", _summary(run["provenance"]["lineage"])))
    if run.get("workload"):
        settings.append(("Workload", run["workload"]))
    sections = [("Test settings and provenance", ["Setting", "Recorded value"], settings),
                ("Measured results", ["Metric", "Value"], [(metric_label(k), metric_value(k, v)) for k, v in sorted(run["metrics"].items())])]
    sections[1:1] = _endpoint_sections(run)
    preflight = run.get("preflight") or {}
    quality = preflight.get("time_quality") or {}
    if preflight:
        clock_rows = [("Preflight status and policy", display(preflight.get("status")) + " | " + display(preflight.get("policy"))),
                      ("Capture timing", display(preflight.get("phase"))),
                      ("Preflight UTC", display(preflight.get("started_at")) + " to " + display(preflight.get("finished_at"))),
                      ("Preflight elapsed seconds", display(preflight.get("elapsed_seconds"))),
                      ("Time quality status", display(quality.get("status"))),
                      ("Clock observation UTC", display(quality.get("observed_at"))),
                      ("Provider and source", display(quality.get("provider")) + " | " + display(quality.get("source"))),
                      ("Provider synchronized", display(quality.get("synchronized")) + " (local provider report; no independent NTP verification)"),
                      ("Reported correction seconds", _clock_seconds(quality.get("offset_seconds"), signed=True)),
                      ("Root dispersion proxy seconds", _clock_seconds(quality.get("uncertainty_seconds")) + " | " + display(quality.get("uncertainty_kind")) + "; not guaranteed clock accuracy"),
                      ("Root delay / dispersion seconds", _clock_seconds(quality.get("root_delay_seconds")) + " / " + _clock_seconds(quality.get("root_dispersion_seconds"))),
                      ("Last synchronization UTC", display(quality.get("last_sync_at"))),
                      ("Clock query elapsed seconds", display(quality.get("query_elapsed_seconds")))]
        sections.append(("Preflight and clock quality", ["Evidence", "Recorded value"], clock_rows))
    else:
        sections.append(("Preflight and clock quality", ["Evidence", "Recorded value"],
                         [("Legacy evidence", "Preflight timing and clock quality were not recorded; no synchronization or accuracy is inferred.")]))
    route = run.get("route") or {}
    if route:
        sections.append(("Route capture context", ["Evidence", "Recorded value"],
                         [("Capture status", display(route.get("status"))),
                          ("Capture phase", display(route.get("phase")) if route.get("phase") else "Timing not recorded in this legacy test"),
                          ("Capture UTC", display(route.get("started_at")) + " to " + display(route.get("finished_at"))),
                          ("Capture elapsed seconds", display(route.get("elapsed_seconds"))),
                          ("Interpretation", "TTL order shows observed replies, not geography or proof of a bottleneck. Missing replies are not measured packet loss. Probe path/source may differ from traffic.")]))
    geography = route.get("geolocation")
    if isinstance(geography, dict):
        build = geography.get("build_epoch")
        try:
            build_date = datetime.fromtimestamp(build, timezone.utc).isoformat() if type(build) is int else "Unavailable"
        except (OverflowError, OSError, ValueError):
            build_date = "Unavailable"
        sections.append(("Approximate location provenance", ["Evidence", "Recorded value"],
                         [("Local dataset status", display(geography.get("status"))),
                          ("Attribution", "IP Geolocation by DB-IP | https://db-ip.com | CC-BY-4.0"),
                          ("Database build UTC", build_date),
                          ("Interpretation", "Approximate public-IP network locations; not GPS, router positions or proof of the physical route. Internal and unknown hops have no invented coordinates.")]))
        location_rows = []
        for hop in route.get("hops", []):
            for alternate, location in enumerate(hop.get("locations", []), 1):
                place = ", ".join(str(location[key]) for key in ("city", "country") if location.get(key))
                coordinates = (f"{location['latitude']:.6g}, {location['longitude']:.6g}"
                               if _finite(location.get("latitude")) and _finite(location.get("longitude")) else "Unavailable")
                location_rows.append((str(hop.get("hop")) + "." + str(alternate), location.get("address") or "No address",
                                      place or "Unknown location", display(location.get("scope")) + " / " + display(location.get("status")), coordinates))
        sections.append(("Approximate route locations", ["TTL / reply", "Observed address", "Approximate place", "Scope / status", "Latitude, longitude"], location_rows))
    diagnostics = run.get("diagnostics") or {}
    latency = diagnostics.get("latency") or {}
    rows = []
    if latency:
        before, loaded = latency.get("baseline") or {}, latency.get("loaded") or {}
        latency_rows = [("Capture status", before.get("status"), loaded.get("status"))]
        keys = list(dict.fromkeys(list((before.get("summary") or {})) + list((loaded.get("summary") or {}))))
        latency_rows.extend((key.replace("_", " "), (before.get("summary") or {}).get(key), (loaded.get("summary") or {}).get(key)) for key in keys)
        sections.append(("ICMP echo RTT under load", ["Metric", "Baseline", "Loaded"], latency_rows))
        rows.append(("Mean increase ms", display(latency.get("mean_increase_ms"))))
    else:
        rows.append(("ICMP echo RTT", "Not captured in this saved test"))
    health = diagnostics.get("endpoint_health") or {}
    if health:
        local = health.get("local") or {}
        rows.append(("Endpoint health scope", "Detected locally; remote metrics unavailable for a plain iPerf peer"))
        rows.append(("Local health delta", _summary(local.get("delta"))))
    else:
        rows.append(("Endpoint health", "Not captured in this saved test"))
    mtu = (diagnostics.get("path_context") or {}).get("mtu") or {}
    rows.extend([("MTU probe status", display(mtu.get("status"))),
                 ("Largest confirmed packet bytes", display(mtu.get("largest_confirmed_packet_bytes"))),
                 ("Route capture status", display((run.get("route") or {}).get("status")))])
    sections.append(("Diagnostic evidence", ["Evidence", "Recorded value"], rows))
    if run.get("flent_summaries"):
        sections.append(("Actual Flent series", ["Series and unit", "Captured summary"],
                         [(str(item.get("name")) + " " + str(item.get("unit")),
                           "; ".join(key + " " + (format_mbps(item.get(key)) if item.get("unit") == "Mbits/s" and key in {"mean", "min", "max"} else display(item.get(key))) for key in ("count", "missing", "mean", "min", "max")))
                          for item in run["flent_summaries"]]))
    return sections


def _route_rows(run):
    rows = []
    for hop in (run.get("route") or {}).get("hops", []):
        observations = [(str(hop.get("hop")), hop)]
        observations.extend((str(hop.get("hop")) + " (obs " + str(index) + ")", observation)
                            for index, observation in enumerate(hop.get("observations", []), 1))
        for label, observation in observations:
            addresses = observation.get("addresses") or ([observation["address"]] if observation.get("address") else [])
            address = ", ".join(addresses) if addresses else "[redacted]" if run.get("redacted") else "No reply or address unavailable"
            values = [display(value) for value in observation.get("rtt_ms", [])]
            values.extend("<" + display(value) for value in observation.get("rtt_upper_bound_ms", []))
            rows.append([label, address, ", ".join(values) if values else "Unavailable"])
    return rows


def _route_diagrams(run):
    """TTL sequence only: arrows do not claim geographic or physical links."""
    from PIL import Image, ImageDraw, ImageFont
    import reportlab
    hops = (run.get("route") or {}).get("hops", [])
    font = ImageFont.truetype(str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"), 26)
    diagrams = []
    for offset in range(0, len(hops), 16):
        panel = hops[offset:offset + 16]
        image = Image.new("RGB", (1500, 360), "white")
        draw = ImageDraw.Draw(image)
        for index, hop in enumerate(panel):
            column, row = index % 8, index // 8
            x, y = 20 + column * 185, 30 + row * 160
            draw.rounded_rectangle((x, y, x + 160, y + 110), radius=8, outline=run["color"], width=3, fill="#f1f4f6")
            draw.text((x + 12, y + 12), "TTL " + str(hop.get("hop")), font=font, fill="#183b50")
            draw.text((x + 12, y + 48), "Reply" if hop.get("responded") else "No reply", font=font, fill="#41494d")
            if index % 8 < 7 and index + 1 < len(panel):
                draw.line((x + 165, y + 55, x + 180, y + 55), fill="#41494d", width=2)
                draw.polygon([(x + 180, y + 55), (x + 174, y + 50), (x + 174, y + 60)], fill="#41494d")
        output = BytesIO()
        image.save(output, "PNG")
        diagrams.append(output.getvalue())
    return diagrams


def _world_map(run):
    """Render bundled public-domain M/L/Z land paths and approximate TTL markers."""
    from PIL import Image, ImageDraw, ImageFont
    import reportlab
    route = run.get("route") or {}
    points = [(hop.get("hop"), index, location) for hop in route.get("hops", [])
              for index, location in enumerate(hop.get("locations", []), 1)
              if location.get("status") == "located" and _finite(location.get("latitude")) and _finite(location.get("longitude"))]
    if not points:
        return None
    image = Image.new("RGB", (1500, 850), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"), 23)
    left, top, scale = 60, 45, 1380 / 900
    draw.rectangle((left, top, left + 1380, top + 690), fill="#eef5fa", outline="#aab6bf", width=2)
    for lon in range(-120, 180, 60):
        x = left + (lon + 180) / 360 * 1380
        draw.line((x, top, x, top + 690), fill="#d9e4ea")
    for lat in range(-60, 90, 30):
        y = top + (90 - lat) / 180 * 690
        draw.line((left, y, left + 1380, y), fill="#d9e4ea")
    asset = Path(__file__).resolve().parent / "static" / "maps" / "world-land.svg"
    for path in ET.parse(asset).getroot().iter("{http://www.w3.org/2000/svg}path"):
        for ring in re.findall(r"M([^MZ]+)Z", path.get("d", "")):
            numbers = [float(number) for number in re.findall(r"-?\d+(?:\.\d+)?", ring)]
            polygon = [(left + x * scale, top + y * scale) for x, y in zip(numbers[::2], numbers[1::2])]
            if len(polygon) >= 3:
                draw.polygon(polygon, fill="#cad6dd", outline="#90a3af")
    for ttl, alternate, location in points:
        x = left + (location["longitude"] + 180) / 360 * 1380
        y = top + (90 - location["latitude"]) / 180 * 690
        draw.ellipse((x - 7, y - 7, x + 7, y + 7), fill=run["color"], outline="white", width=2)
        label = str(ttl) + "." + str(alternate)
        # Latitude/longitude remain in the table; clamp labels inside the map.
        draw.text((min(left + 1290, max(left + 5, x + 12)), min(top + 650, max(top + 5, y - 27))),
                  label, font=font, fill="#183b50", stroke_width=2, stroke_fill="white")
    draw.text((60, 760), "Markers: TTL.alternate reply; no physical route is inferred.", font=font, fill="#41494d")
    draw.text((60, 800), "Natural Earth land: public domain. Location attribution is listed in this report.", font=font, fill="#41494d")
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def _native_chart(runs, unit, bounds=None):
    """Plot Flent samples as points, keeping unknown samples and distinct units apart."""
    from PIL import Image, ImageDraw, ImageFont
    import reportlab
    font = ImageFont.truetype(str(Path(reportlab.__file__).parent / "fonts" / "Vera.ttf"), 21)
    groups = [(run_index, item) for run_index, run in enumerate(runs) for item in run.get("flent_series", []) if item.get("unit") == unit]
    image = Image.new("RGB", (1500, 620), "white")
    draw = ImageDraw.Draw(image)
    points = [point for _, group in groups for point in group.get("points", []) if _finite(point.get("seconds")) and _finite(point.get("value")) and point["seconds"] >= 0 and point["value"] >= 0]
    if not points:
        draw.text((100, 200), "No captured samples for this unit", font=font, fill="#41494d")
    else:
        bounds = bounds or _bounds(runs, unit)
        xmax, ymax, scale = bounds["xmax"], bounds["ymax"], bounds["scale"]
        draw.text((15, 20), scale["unit"], font=font, fill="#1b1f22")
        for tick in range(5):
            y, x = 490 - 390 * tick / 4, 130 + 1300 * tick / 4
            draw.line((130, y, 1430, y), fill="#d9d9d9")
            draw.text((15, y - 12), f"{ymax * tick / 4:.1f}", font=font, fill="#41494d")
            draw.text((x - 20, 510), f"{xmax * tick / 4:.1f}", font=font, fill="#41494d")
        draw.text((570, 560), "Elapsed seconds", font=font, fill="#41494d")
        for group_index, (run_index, group) in enumerate(groups):
            color = runs[run_index]["color"]
            for point in group.get("points", []):
                if not _finite(point.get("seconds")) or not _finite(point.get("value")) or point["seconds"] < 0 or point["value"] < 0:
                    continue
                x, y = 130 + point["seconds"] / xmax * 1300, 490 - (point["value"] * bounds["multiplier"] / scale["divisor"]) / ymax * 390
                if group_index % 2:
                    draw.rectangle((x - 4, y - 4, x + 4, y + 4), outline=color, width=2)
                else:
                    draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill=color)
        draw.text((130, 585), "Named series and saved colors are listed below", font=font, fill="#41494d")
    output = BytesIO()
    image.save(output, "PNG")
    return output.getvalue()


def _charts(runs):
    if any(run.get("kind") == "diagnostic-check" for run in runs):
        return []
    native = [run for run in runs if run.get("flent_series")]
    charts = []
    if len(native) == len(runs):
        for unit in ("Mbits/s", "ms"):
            if any(item.get("unit") == unit for run in native for item in run["flent_series"]):
                for offset in range(0, len(runs), 5):
                    panel = runs[offset:offset + 5]
                    groups = [(run, item) for run in panel for item in run["flent_series"] if item.get("unit") == unit]
                    legend = [("Test " + str(run["display_index"]), run["color"], str(item.get("name")) + (" (outlined square)" if index % 2 else " (circle)"))
                              for index, (run, item) in enumerate(groups)]
                    charts.append(("Actual Flent throughput samples" if unit == "Mbits/s" else "Actual Flent round trip samples", _native_chart(panel, unit, _bounds(runs, unit)),
                                   "Named native Flent series are individual samples; missing samples stay blank. SI throughput units and RTT remain separate. Panels share axes. Markers alternate circles and outlined squares; colors can repeat.", legend))
    elif not native:
        bounds = _bounds(runs)
        for offset in range(0, len(runs), 10):
            panel = runs[offset:offset + 10]
            legend = [("Test " + str(run["display_index"]), run["color"], run["label"]) for run in panel]
            custom_direction = any(point.get("direction") not in KNOWN_DIRECTIONS for run in panel for point in run["series"])
            caption = ("Known upload intervals are solid; download dashed; other directions are points. Only captured data are drawn, with shared axes across panels. Short traces stop at their last measurement. Saved colors can repeat."
                       if custom_direction else "Upload solid; download dashed. Only captured spans and points are drawn, with shared axes across panels. Short traces stop at their last measurement. Trace labels identify saved colors, which can repeat.")
            charts.append(("Measured throughput intervals", _chart(panel, bounds),
                           caption, legend))
    return charts


def _context(run):
    rows = []
    if run.get("project"):
        project = run["project"]
        if isinstance(project, dict):
            rows.append("Project: " + str(project.get("name", "")))
            if project.get("description"):
                rows.append("Project description: " + str(project["description"]))
            if project.get("baseline_id"):
                rows.append("Project baseline ID: " + project["baseline_id"])
        else:
            rows.append("Project: " + display(project))
    if run.get("tags"):
        rows.append("Tags: " + display(run["tags"]))
    if run.get("notes"):
        rows.append("Notes: " + run["notes"])
    for note in run.get("annotations", []):
        rows.append(f"{note.get('timestamp')} | elapsed {display(note.get('seconds'))} s | {note.get('text')}")
    return rows


def _window_rows(model):
    rows = []
    window = model["common_window"]
    for index, run in enumerate(window["runs"], 1):
        values = []
        for direction in run.get("directions", []):
            values.append(str(direction.get("direction")) + ": " + (format_mbps(direction["mean_mbps"]) if _finite(direction.get("mean_mbps")) else "Unavailable")
                          + (" (" + direction["reason"] + ")" if direction.get("reason") else ""))
        rows.append(("Test " + str(index), display(run.get("measured_end_seconds")) + " s", "; ".join(values) or "Unavailable"))
    return rows


def render_pdf(model):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfbase.pdfdoc import PDFDictionary, PDFArray, PDFStream, PDFString, PDFName
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
    import reportlab
    with FONT_LOCK:
        for name, filename in (("FlowBenchVera", "Vera.ttf"), ("FlowBenchVeraBold", "VeraBd.ttf")):
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont(name, str(Path(reportlab.__file__).parent / "fonts" / filename)))
        pdfmetrics.registerFontFamily("FlowBenchVera", normal="FlowBenchVera", bold="FlowBenchVeraBold")
    allowed = pdfmetrics.getFont("FlowBenchVera").face.charToGlyph
    def safe(text):
        return escape("".join(c if ord(c) in allowed or c in "\n\t" else "?" for c in _clean(str(text)))).replace("\n", "<br/>")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("FBBody", fontName="FlowBenchVera", fontSize=9, leading=13, spaceAfter=5, splitLongWords=True))
    styles.add(ParagraphStyle("FBHeading", fontName="FlowBenchVeraBold", fontSize=13, leading=17, spaceBefore=13, spaceAfter=7, keepWithNext=True))
    styles.add(ParagraphStyle("FBTitle", fontName="FlowBenchVeraBold", fontSize=23, leading=28, spaceAfter=12))
    styles.add(ParagraphStyle("FBCell", fontName="FlowBenchVera", fontSize=8, leading=11))
    def p(text, style="FBBody"):
        return Paragraph(safe(text), styles[style])
    image_bytes = 0
    def chart_image(payload, height=205):
        nonlocal image_bytes
        image_bytes += len(payload)
        if image_bytes > MODEL_BYTES:
            raise ValueError("Report exceeds the 32 MiB chart-image resource budget. Split the selection into smaller reports.")
        return Image(BytesIO(payload), width=495, height=height)
    styles.add(ParagraphStyle("FBHeaderCell", parent=styles["FBCell"], textColor=colors.white, fontName="FlowBenchVeraBold"))
    def table(head, rows, widths=None):
        widths = widths or ((60, 145, 70, 45, 175) if len(head) == 5 and head[0] == "Section" else (60, 130, 100, 105, 100) if len(head) == 5 and head[0] == "Attempt" else (45, 120, 140, 90, 100) if len(head) == 5 else (210, 140, 145) if len(head) == 3 else (170, 325))
        content = [[p(item, "FBHeaderCell") for item in head]] + [[p(display(item), "FBCell") for item in row] for row in rows]
        result = Table(content, colWidths=widths, repeatRows=1, hAlign="LEFT")
        result.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#183b50")),
                                   ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                   ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f4f6")]),
                                   ("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#d9d9d9")),
                                   ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                                   ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
        return result
    flow = [Image(str(BRAND), width=180, height=60, hAlign="LEFT"), Spacer(1, 14), p(model["title"], "FBTitle"),
            p(f"{len(model['runs'])} {'terminal check(s)' if model.get('report_type') == 'diagnostic-check' else 'completed test(s)'} | Generated {model['generated_at']} | FlowBench {model['version']}"),
            p("This report records captured performance and diagnostic evidence. The first selected test is the baseline for any comparisons."),
            p("Redaction " + ("enabled" if model["redacted"] else "disabled") + "; raw evidence " + ("requested" if model["include_raw"] else "excluded") + ".")]
    if len(model["runs"]) > 1 and model.get("common_window"):
        flow.extend([p("Compared elapsed window", "FBHeading"), p(model["common_window"]["note"]),
                     p("Common captured window: " + display(model["common_window"].get("seconds")) + " seconds. Summary comparisons below retain whole-run aggregates."),
                     table(["Test", "Measured end", "Common window mean"], _window_rows(model), (65, 95, 335))])
    for report in model["comparisons"]:
        selected = next(row["display_index"] for row in model["runs"] if row["id"] == report["after_id"])
        flow.append(KeepTogether([p("Baseline compared with " + ("Check " if model.get("report_type") == "diagnostic-check" else "Test ") + str(selected), "FBHeading"), p(report["caveat"]),
                     p("Different settings: " + (", ".join(report["settings_differences"]) or "None detected")),
                     table(["Metric", "Baseline", "Selected", "Change" if model.get("report_type") == "diagnostic-check" else "Change %"],
                           [[metric_label(row["metric"]), metric_value(row["metric"], row["before"]), metric_value(row["metric"], row["after"]), metric_value(row["metric"], row.get("delta")) if model.get("report_type") == "diagnostic-check" else display(row["percent_change"])] for row in report["metrics"]], (210, 95, 95, 95))]))
        if report.get("details"):
            flow.append(table(["Evidence", "Recorded value"], report["details"]))
    flow.append(p("Interpretation and limitations", "FBHeading"))
    flow.extend(p(text) for text in model["limitations"])
    if len(model["runs"]) > 1 and model.get("report_type") != "diagnostic-check":
        flow.append(PageBreak())
        charts = _charts(model["runs"])
        for title, chart, caption, legend in charts:
            flow.append(KeepTogether([p(title, "FBHeading"), chart_image(chart), p(caption)]))
            flow.append(table(["Trace", "Saved color", "Label or native series"], legend, (65, 95, 335)))
        if not charts:
            flow.append(p("These workloads have different measurement series; inspect their individual charts and comparison caveats below."))
    for index, run in enumerate(model["runs"], 1):
        flow.extend([PageBreak(), p(f"{'Check' if run.get('kind') == 'diagnostic-check' else 'Test'} {index}", "FBTitle"), p(run["label"])])
        for heading, head, rows in _sections(run):
            flow.extend([p(heading, "FBHeading"), table(head, rows)])
            if heading == "Route capture context":
                for diagram in _route_diagrams(run):
                    flow.extend([chart_image(diagram, height=119), p("Observed TTL sequence. Arrows show ordering only; reply addresses and measured RTTs appear in the hop table.")])
            if heading == "Approximate route locations":
                world = _world_map(run)
                if world:
                    flow.append(KeepTogether([p("Approximate public-IP world map", "FBHeading"), chart_image(world, height=280),
                                              p("Approximate network locations only. Missing and internal hops stay unlocated; markers do not prove router positions or the physical route. Coincident markers can overlap; the table retains every reply.")]))
            if heading == "Measured results":
                for title, chart, caption, legend in _charts([run]):
                    flow.append(KeepTogether([p(title, "FBHeading"), chart_image(chart), p(caption)]))
                    flow.append(table(["Trace", "Saved color", "Label or native series"], legend, (65, 95, 335)))
        context = _context(run)
        if context:
            flow.append(p("Project labels and timeline notes", "FBHeading"))
            flow.extend(p(text) for text in context)
        if isinstance(run.get("route"), dict) and run["route"].get("hops"):
            flow.extend([p("Captured route hops", "FBHeading"),
                         table(["Hop", "Address or response", "RTT ms"], _route_rows(dict(run, redacted=model["redacted"])), (50, 290, 155))])
        if run.get("raw_evidence"):
            flow.extend([PageBreak(), p("Raw evidence", "FBHeading"), p("Raw fields are limited to 20000 characters each. Any truncation is explicitly recorded below."), p(display(run["raw_evidence"]))])
    output = BytesIO()
    def footer(canvas, doc):
        if doc.page > REPORT_PAGES:
            raise ValueError("Report exceeds the 500 PDF-page resource budget. Split the selection into smaller reports.")
        if doc.page == 1:
            license_bytes = (Path(__file__).resolve().parent / "third_party" / "licenses" / "bitstream-vera-LICENSE.txt").read_bytes()
            embedded = canvas._doc.Reference(PDFStream(PDFDictionary({"Type": PDFName("EmbeddedFile")}), license_bytes, filters=[]))
            specification = canvas._doc.Reference(PDFDictionary({"Type": PDFName("Filespec"), "F": PDFString("bitstream-vera-LICENSE.txt"),
                                                                 "EF": PDFDictionary({"F": embedded})}))
            canvas._doc.Catalog.Names = PDFDictionary({"EmbeddedFiles": PDFDictionary({"Names": PDFArray([PDFString("bitstream-vera-LICENSE.txt"), specification])})})
        canvas.saveState()
        canvas.setFont("FlowBenchVera", 8)
        canvas.setFillColor(colors.HexColor("#41494d"))
        canvas.drawString(50, 30, "FlowBench | Captured network measurements")
        canvas.drawRightString(A4[0] - 50, 30, f"Page {doc.page}")
        canvas.restoreState()
    SimpleDocTemplate(output, pagesize=A4, leftMargin=50, rightMargin=50, topMargin=44, bottomMargin=50,
                      title=model["title"], author="FlowBench", pageCompression=1).build(flow, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def render_docx(model):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.7)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Title"].font.size = Pt(25)
    document.styles["Heading 1"].font.size = Pt(15)
    document.styles["Heading 2"].font.size = Pt(12)
    # The bundled Word Title style can inherit a themed blue bottom rule.
    # Remove paragraph borders from styles, keeping explicit table borders.
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    document.core_properties.title = model["title"]
    document.core_properties.author = "FlowBench"
    document.core_properties.subject = "Captured network measurements"
    header = section.header.paragraphs[0]
    header.text = "FlowBench"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("FlowBench | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    image_bytes = 0
    def add_chart(payload, height=2.8):
        nonlocal image_bytes
        image_bytes += len(payload)
        if image_bytes > MODEL_BYTES:
            raise ValueError("Report exceeds the 32 MiB chart-image resource budget. Split the selection into smaller reports.")
        document.add_picture(BytesIO(payload), width=Inches(6.8), height=Inches(height))
        document.paragraphs[-1].paragraph_format.keep_with_next = True
    def para(text):
        document.add_paragraph(_clean(str(text)))
    def table(head, rows):
        result = document.add_table(rows=1, cols=len(head))
        result.autofit = False
        widths = [.8, 1.8, .9, .5, 2.8] if len(head) == 5 and head[0] == "Section" else [.8, 1.9, 1.3, 1.4, 1.4] if len(head) == 5 and head[0] == "Attempt" else [.6, 1.65, 1.85, 1.3, 1.4] if len(head) == 5 else [3.25, 1.2, 1.2, 1.15] if len(head) == 4 else [.8, 4.4, 1.6] if len(head) == 3 and head[0] == "Hop" else [0.8, 1.2, 4.8] if len(head) == 3 and head[0] in {"Test", "Trace"} else [2.8, 2, 2] if len(head) == 3 else [2.2, 4.6]
        for column, width in zip(result.columns, widths):
            column.width = Inches(width)
        for cell, title in zip(result.rows[0].cells, head):
            cell.text = title
        for row in rows:
            for cell, text in zip(result.add_row().cells, row):
                cell.text = _clean(display(text))
        repeat = OxmlElement("w:tblHeader")
        result.rows[0]._tr.get_or_add_trPr().append(repeat)
        for index, row in enumerate(result.rows):
            # Keep each row intact, but allow the table itself to paginate.
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement("w:cantSplit"))
            for column, cell in enumerate(row.cells):
                cell.width = Inches(widths[column])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "183B50" if index == 0 else "F1F4F6" if index % 2 == 0 else "FFFFFF")
                props.append(shade)
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement("w:" + side)
                    for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                        edge.set(qn("w:" + attr), value)
                    borders.append(edge)
                props.append(borders)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement("w:" + side)
                    edge.set(qn("w:w"), "100")
                    edge.set(qn("w:type"), "dxa")
                    margins.append(edge)
                props.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2)
                    # LibreOffice can orphan a repeated table header at a page
                    # boundary and omit its first body row during PDF conversion.
                    p.paragraph_format.keep_with_next = index == 0
                    if len(head) == 4 and column > 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
        para("")
    document.add_picture(str(BRAND), width=Inches(2.4))
    document.add_paragraph(model["title"], "Title")
    para(f"{len(model['runs'])} {'terminal check(s)' if model.get('report_type') == 'diagnostic-check' else 'completed test(s)'} | Generated {model['generated_at']} | FlowBench {model['version']}")
    para("This report records captured performance and diagnostic evidence. The first selected test is the baseline for any comparisons.")
    para("Redaction " + ("enabled" if model["redacted"] else "disabled") + "; raw evidence " + ("requested" if model["include_raw"] else "excluded") + ".")
    if len(model["runs"]) > 1 and model.get("common_window"):
        document.add_heading("Compared elapsed window", 1)
        para(model["common_window"]["note"])
        para("Common captured window: " + display(model["common_window"].get("seconds")) + " seconds. Summary comparisons below retain whole-run aggregates.")
        table(["Test", "Measured end", "Common window mean"], _window_rows(model))
    for report in model["comparisons"]:
        selected = next(row["display_index"] for row in model["runs"] if row["id"] == report["after_id"])
        heading = document.add_heading("Baseline compared with " + ("Check " if model.get("report_type") == "diagnostic-check" else "Test ") + str(selected), 1)
        # Native LibreOffice can strand a comparison heading even with direct
        # keepNext properties. Begin the bounded introductory block on a fresh
        # page; the table can still paginate after its first measured row.
        heading.paragraph_format.page_break_before = True
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.keep_together = True
        para(report["caveat"])
        document.paragraphs[-1].paragraph_format.keep_with_next = True
        para("Different settings: " + (", ".join(report["settings_differences"]) or "None detected"))
        document.paragraphs[-1].paragraph_format.keep_with_next = True
        table(["Metric", "Baseline", "Selected", "Change" if model.get("report_type") == "diagnostic-check" else "Change %"], [[metric_label(r["metric"]), metric_value(r["metric"], r["before"]), metric_value(r["metric"], r["after"]), metric_value(r["metric"], r.get("delta")) if model.get("report_type") == "diagnostic-check" else r["percent_change"]] for r in report["metrics"]])
        if report.get("details"):
            table(["Evidence", "Recorded value"], report["details"])
    document.add_heading("Interpretation and limitations", 1)
    for text in model["limitations"]:
        para(text)
    if len(model["runs"]) > 1 and model.get("report_type") != "diagnostic-check":
        document.add_page_break()
        charts = _charts(model["runs"])
        for title, chart, caption, legend in charts:
            document.add_heading(title, 1)
            add_chart(chart)
            para(caption)
            table(["Trace", "Saved color", "Label or native series"], legend)
        if not charts:
            para("These workloads have different measurement series; inspect their individual charts and comparison caveats below.")
    for index, run in enumerate(model["runs"], 1):
        heading = document.add_heading(f"{'Check' if run.get('kind') == 'diagnostic-check' else 'Test'} {index}", 1)
        # A standalone break paragraph can itself move onto a fresh page and
        # create an empty sheet after a full evidence table in LibreOffice.
        heading.paragraph_format.page_break_before = True
        heading.paragraph_format.space_before = Pt(0)
        para(run["label"])
        for heading, head, rows in _sections(run):
            document.add_heading(heading, 2)
            table(head, rows)
            if heading == "Route capture context":
                for diagram in _route_diagrams(run):
                    add_chart(diagram, height=1.63)
                    para("Observed TTL sequence. Arrows show ordering only; reply addresses and measured RTTs appear in the hop table.")
            if heading == "Approximate route locations":
                world = _world_map(run)
                if world:
                    document.add_heading("Approximate public-IP world map", 2)
                    add_chart(world, height=3.85)
                    para("Approximate network locations only. Missing and internal hops stay unlocated; markers do not prove router positions or the physical route. Coincident markers can overlap; the table retains every reply.")
            if heading == "Measured results":
                for title, chart, caption, legend in _charts([run]):
                    document.add_heading(title, 2)
                    add_chart(chart)
                    para(caption)
                    table(["Trace", "Saved color", "Label or native series"], legend)
        if _context(run):
            document.add_heading("Project labels and timeline notes", 2)
            for text in _context(run):
                para(text)
        if isinstance(run.get("route"), dict) and run["route"].get("hops"):
            document.add_heading("Captured route hops", 2)
            table(["Hop", "Address or response", "RTT ms"], _route_rows(dict(run, redacted=model["redacted"])))
        if run.get("raw_evidence"):
            document.add_page_break()
            document.add_heading("Raw evidence", 1)
            para("Raw fields are limited to 20000 characters each. Any truncation is explicitly recorded below.")
            para(display(run["raw_evidence"]))
    output = BytesIO()
    # Table spacing is useful between sections, but a final empty paragraph can
    # overflow onto an otherwise blank page in native LibreOffice conversion.
    # Body-level tables do not require the trailing paragraph that cells do.
    body = document._element.body
    while len(body) > 1 and body[-2].tag == qn("w:p") and len(body[-2]) == 0:
        body.remove(body[-2])
    for border in document._element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    content_bytes = len(document._element.xml.encode("utf-8")) + sum(len(part.blob) for part in document.part.package.image_parts)
    if content_bytes > MODEL_BYTES:
        raise ValueError("Report exceeds the 32 MiB Word-content resource budget. Split the selection into smaller reports.")
    document.save(output)
    return output.getvalue()
