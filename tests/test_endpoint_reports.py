"""Endpoint reports separate native sockets from sensitive HTTP provenance."""
import copy
from io import BytesIO
import json
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from docx.oxml.ns import qn

from app.endpoints import endpoint_view
from app.exports import xml_export
from app.imports import import_runs, parse_document
from app.reports import prepare_report, render_pdf, render_docx, _sections, _report_endpoints
from app.runner import RunManager
try:
    from test_endpoint_imports import endpoint_fixture
    from test_reports import report_fixture, decoded_pdf_streams
except ModuleNotFoundError:
    from tests.test_endpoint_imports import endpoint_fixture
    from tests.test_reports import report_fixture, decoded_pdf_streams


class EndpointReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.manager = RunManager(self.temp.name, "missing-iperf")

    def tearDown(self):
        self.manager.shutdown()
        self.temp.cleanup()

    def test_imported_json_xml_default_redaction_hides_all_endpoint_context_even_raw(self):
        for xml in (False, True):
            source = endpoint_fixture()
            source["id"] = ("b" if xml else "c") * 32
            payload = xml_export(source) if xml else json.dumps(source).encode()
            identity = import_runs(self.manager, parse_document(payload), "fixture")["imported"][0]
            before = self.manager.export(identity)
            model = prepare_report(self.manager, {"ids": [identity], "include_raw": True})
            endpoint = model["runs"][0]["endpoints"]
            self.assertEqual(endpoint["socket_count"], 2)
            self.assertEqual(endpoint["families"], [4, 6])
            self.assertEqual(endpoint["provenance"], "imported_archive")
            outputs = [json.dumps(model).encode(), decoded_pdf_streams(render_pdf(model))]
            with ZipFile(BytesIO(render_docx(model))) as package:
                outputs.append(b"\n".join(package.read(name) for name in package.namelist() if name.endswith(".xml")))
            for output in outputs:
                for private in (b"execution-private", b"private_scope", b"192.0.2.211", b"198.51.100.213", b"fe80::211", b"2001:db8::213", b"203.0.113.217", b"43117", b"43118"):
                    self.assertNotIn(private, output)
            self.assertEqual(before, self.manager.export(identity))

    def test_unknown_numeric_fields_and_nested_copies_never_enter_redacted_projection(self):
        source = endpoint_fixture()
        # Simulate future helper additions without widening the report contract.
        view = endpoint_view(source)
        view.update(future_numeric_ip=3405803985, future_peer={"address": 3405803985, "port": 49199})
        with patch("app.endpoints.endpoint_view", return_value=view):
            projected, context = _report_endpoints(source, True)
        self.assertNotIn("3405803985", json.dumps(projected))
        self.assertNotIn("49199", json.dumps(projected))
        # Historical archive copies may exist beneath generic diagnostic keys.
        source["diagnostics"]["latency"]["before"] = {"endpoints": copy.deepcopy(source["endpoints"]),
            "initiator_context": copy.deepcopy(source["initiator_context"]), "unknown_numeric_ip": 3405803985,
            "alias": dict(source["initiator_context"], secret_number=3405803985)}
        self.manager.save(source)
        model = prepare_report(self.manager, {"ids": [source["id"]], "include_raw": True})
        outputs = [json.dumps(model).encode(), decoded_pdf_streams(render_pdf(model))]
        with ZipFile(BytesIO(render_docx(model))) as package:
            outputs.append(b"\n".join(package.read(name) for name in package.namelist() if name.endswith(".xml")))
        for output in outputs:
            for secret in (b"3405803985", b"203.0.113.217", b"43117"):
                self.assertNotIn(secret, output)

    def test_unredacted_human_sections_keep_perspective_and_fingerprint_tracks_context(self):
        record = endpoint_fixture()
        self.manager.save(record)
        original = self.manager.export(record["id"])
        model = prepare_report(self.manager, {"ids": [record["id"]], "redact": False})
        text = str(_sections(model["runs"][0]))
        for expected in ("execution-private.example", "Local 192.0.2.211:43117", "Remote 198.51.100.213:5201", "[fe80::211%private_scope]:43118", "Initiating HTTP peer (not traffic source)", "203.0.113.217", "NAT egress", "bidirectional"):
            self.assertIn(expected, text)
        self.assertEqual(model["runs"][0]["endpoints"]["sockets"], record["endpoints"]["sockets"])
        self.assertEqual(original, self.manager.export(record["id"]))
        record["initiator_context"]["peer"]["address"] = "203.0.113.218"
        self.manager.save(record)
        with self.assertRaisesRegex(ValueError, "changed"):
            prepare_report(self.manager, {"ids": [record["id"]], "redact": False, "fingerprint": model["fingerprint"]})

    def test_legacy_view_is_pure_and_invalid_historical_metadata_rejected_without_rewrite(self):
        record = endpoint_fixture()
        record.pop("endpoints")
        record.pop("initiator_context")
        self.manager.save(record)
        before = self.manager.export(record["id"])
        with patch("app.endpoints.runtime_hostname", side_effect=AssertionError("No identity lookup")), patch("socket.getaddrinfo", side_effect=AssertionError("No network")):
            model = prepare_report(self.manager, {"ids": [record["id"]], "redact": False})
        self.assertEqual(model["runs"][0]["endpoints"]["provenance"], "preserved_native")
        self.assertIsNone(model["runs"][0]["endpoints"]["execution_hostname"])
        self.assertEqual(before, self.manager.export(record["id"]))

        old = report_fixture()
        old["id"] = "d" * 32
        self.manager.save(old)
        absent = prepare_report(self.manager, {"ids": [old["id"]], "redact": False})
        self.assertEqual(absent["runs"][0]["endpoints"]["status"], "not_recorded")
        self.assertNotIn("endpoints", self.manager.read(old["id"]))
        record["initiator_context"] = {"schema_version": 99}
        self.manager.save(record)
        before = self.manager.export(record["id"])
        with self.assertRaises(ValueError):
            prepare_report(self.manager, {"ids": [record["id"]]})
        self.assertEqual(before, self.manager.export(record["id"]))

    def test_word_final_table_does_not_create_an_empty_spacer_page(self):
        record = endpoint_fixture()
        self.manager.save(record)
        model = prepare_report(self.manager, {"ids": [record["id"]], "redact": False})
        document = Document(BytesIO(render_docx(model)))
        body = document._element.body
        self.assertEqual(body[-1].tag, qn("w:sectPr"))
        self.assertEqual(body[-2].tag, qn("w:tbl"))
        # The complete final observation remains, with its measured RTTs.
        self.assertEqual(document.tables[-1].cell(1, 1).text, "192.0.2.1")
        self.assertEqual(document.tables[-1].cell(1, 2).text, "1.2, 1.3")


if __name__ == "__main__":
    unittest.main()
