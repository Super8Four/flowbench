"""Standalone diagnostic reports retain failures without inventing throughput."""
import copy
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document

from app.checks import summarize_check, check_run_status
from app.exports import xml_export
from app.imports import import_runs, parse_document
from app.reports import prepare_report, render_pdf, render_docx, _private_check, _sections
from app.runner import RunManager
try:
    from test_reports import decoded_pdf_streams, report_fixture
except ModuleNotFoundError:
    from tests.test_reports import decoded_pdf_streams, report_fixture


spec = importlib.util.spec_from_file_location("report_fixtures", Path(__file__).resolve().parents[1] / "scripts" / "check-reports.py")
fixtures = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixtures)


class CheckReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.manager = RunManager(self.temp.name, "missing-iperf")

    def tearDown(self):
        self.manager.shutdown()
        self.temp.cleanup()

    def save(self, kind, identity="e" * 32, partial=False):
        record = fixtures.check_fixture(kind, identity, partial)
        self.manager.save(record)
        return record

    def test_typed_terminal_checks_never_acquire_iperf_settings_or_metrics(self):
        for kind in ("dns", "loss"):
            for partial in (False, True):
                record = self.save(kind, partial=partial)
                model = prepare_report(self.manager, {"ids": [record["id"]]})
                run = model["runs"][0]
                self.assertEqual(run["metrics"], {})
                self.assertEqual(run["series"], [])
                self.assertEqual(run["status"], "partial" if partial else "completed")
                self.assertNotIn("protocol", run["settings"])
                self.assertNotIn("streams", run["settings"])
                self.assertIsNone(model["common_window"])
                titles = [section[0] for section in _sections(run)]
                self.assertNotIn("Measured results", titles)
                self.assertIn("Captured diagnostic summary", titles)
                self.assertIn("Clock quality", titles)

    def test_dns_comparison_redacts_names_answers_raw_and_freeform_provenance(self):
        before = self.save("dns")
        before["result"]["adapter"].update(version="private-version.example", platform="private-platform.example")
        before["result"]["note"] = "private-note.example"
        rr = before["result"]["attempts"][1]["answers"][1]
        rr.update(type="private-type.example", **{"class": "private-class.example"}, rdata='"private-answer.example 12.345 67.891"')
        before["result"]["attempts"][0]["raw_output"] = "private-raw.example"
        self.manager.save(before)
        after = self.save("dns", "f" * 32, partial=True)
        original = self.manager.export(before["id"])
        model = prepare_report(self.manager, {"ids": [before["id"], after["id"]], "include_raw": True})
        self.assertEqual(len(model["comparisons"]), 1)
        self.assertNotIn("answer_sets", model["comparisons"][0])
        sensitive = ("report-fixture.example", "192.0.2.53", "192.0.2.25", "private-version.example", "private-platform.example", "private-note.example", "private-answer.example", "private-type.example", "private-class.example", "private-raw.example", "12.345", "67.891")
        pdf = decoded_pdf_streams(render_pdf(model))
        with ZipFile(BytesIO(render_docx(model))) as package:
            xml = b"\n".join(package.read(name) for name in package.namelist() if name.endswith(".xml"))
        for secret in sensitive:
            self.assertNotIn(secret, json.dumps(model))
            self.assertNotIn(secret.encode(), pdf)
            self.assertNotIn(secret.encode(), xml)
        self.assertEqual(model["runs"][0]["result"]["attempts"][1]["answers"][1]["ttl"], 0)
        self.assertEqual(model["runs"][0]["result"]["attempts"][2]["status"], "nodata")
        self.assertEqual(model["runs"][0]["result"]["attempts"][3]["rcode"], "NXDOMAIN")
        self.assertEqual(self.manager.export(before["id"]), original)

    def test_unknown_numeric_fields_are_denied_and_invalid_archive_evidence_rejected(self):
        private = _private_check({"summary": {"rtt": {"mean_ms": 2.5, "private_coordinate": 12.345}}, "private_numeric": 67.891})
        self.assertEqual(private, {"summary": {"rtt": {"mean_ms": 2.5}}})
        record = self.save("loss")
        record["result"]["private_numeric"] = 67.891
        self.manager.save(record)
        original = self.manager.export(record["id"])
        with self.assertRaises(ValueError):
            prepare_report(self.manager, {"ids": [record["id"]]})
        self.assertEqual(self.manager.export(record["id"]), original)

    def test_icmp_unknown_denominator_and_rtt_bounds_stay_missing(self):
        record = self.save("loss", partial=True)
        model = prepare_report(self.manager, {"ids": [record["id"]]})
        result = model["runs"][0]["result"]
        self.assertIsNone(result["summary"]["transmitted"])
        self.assertIsNone(result["summary"]["nonresponse_percent"])
        self.assertEqual(result["summary"]["rtt"]["count"], 1)
        self.assertEqual(result["summary"]["rtt"]["mean_ms"], 2.5)
        self.assertIsNone(result["attempts"][1]["rtt_ms"])
        self.assertEqual(result["attempts"][1]["rtt_upper_bound_ms"], 1)
        doc = Document(BytesIO(render_docx(model)))
        text = " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
        self.assertIn("<1", text)
        self.assertIn("permission_blocked", text)
        self.assertIn("Unavailable", text)
        self.assertNotIn("Mbps", text)
        record["result"]["attempts"][0]["rtt_ms"] = .0000004
        record["result"]["summary"] = summarize_check("loss", record["result"], record["settings"])
        self.manager.save(record)
        tiny = prepare_report(self.manager, {"ids": [record["id"]]})
        probes = next(rows for title, _, rows in _sections(tiny["runs"][0]) if title == "ICMP echo observations")
        self.assertEqual(probes[0][2], "4e-07")
        self.assertEqual(tiny["runs"][0]["result"]["summary"]["rtt"]["mean_ms"], .0000004)

    def test_failed_cancelled_interrupted_checks_report_honest_status_and_reject_active(self):
        for state in ("error", "unavailable", "cancelled", "interrupted", "running", "pending"):
            record = fixtures.check_fixture("dns", partial=True)
            record["result"].update(status=state, attempts=[])
            if state == "pending":
                record["result"]["started_at"] = None
            record["result"]["summary"] = summarize_check("dns", record["result"], record["settings"])
            record["status"] = check_run_status(record["result"])
            self.manager.save(record)
            if state in {"pending", "running"}:
                with self.assertRaises(ValueError):
                    prepare_report(self.manager, {"ids": [record["id"]]})
            else:
                model = prepare_report(self.manager, {"ids": [record["id"]]})
                self.assertEqual(model["runs"][0]["status"], record["status"])
                self.assertEqual(model["runs"][0]["result"]["summary"]["responses"], 0)

    def test_unredacted_dns_literal_text_and_raw_opt_in_are_separate(self):
        record = self.save("dns")
        model = prepare_report(self.manager, {"ids": [record["id"]], "redact": False})
        self.assertNotIn("raw_output", json.dumps(model))
        self.assertIn("Synthetic <literal>& record", json.dumps(model))
        with ZipFile(BytesIO(render_docx(model))) as package:
            self.assertIn(b"&lt;literal&gt;&amp;", package.read("word/document.xml"))
        raw = prepare_report(self.manager, {"ids": [record["id"]], "redact": False, "include_raw": True})
        self.assertIn("192.0.2.53", raw["runs"][0]["raw_evidence"]["native_result"])
        self.assertNotEqual(model["fingerprint"], raw["fingerprint"])

    def test_mixed_checks_or_throughput_selections_are_rejected(self):
        dns, loss = self.save("dns"), self.save("loss", "f" * 32)
        throughput = report_fixture()
        self.manager.save(throughput)
        for other in (loss, throughput):
            with self.assertRaises(ValueError):
                prepare_report(self.manager, {"ids": [dns["id"], other["id"]]})

    def test_lineage_and_absolute_comparison_changes_retain_their_meaning(self):
        before = self.save("loss")
        after = self.save("loss", "f" * 32, partial=True)
        after["lineage"] = {"schema_version": 1, "action": "run_again", "source_run_id": before["id"], "external_source_id": None}
        self.manager.save(after)
        model = prepare_report(self.manager, {"ids": [before["id"], after["id"]]})
        self.assertEqual(model["runs"][1]["provenance"]["lineage"], after["lineage"])
        document = Document(BytesIO(render_docx(model)))
        run_headings = [paragraph for paragraph in document.paragraphs
                        if paragraph.style.name == "Heading 1" and paragraph.text in ("Check 1", "Check 2")]
        self.assertEqual(len(run_headings), 2)
        self.assertTrue(all(paragraph.paragraph_format.page_break_before for paragraph in run_headings))
        # A separate break paragraph can overflow and manufacture a blank page.
        self.assertEqual(document._element.xpath('.//w:br[@w:type="page"]'), [])
        for row in model["comparisons"][0]["metrics"]:
            self.assertNotIn("percent_change", row)
            self.assertIn("delta", row)
        percentage = next(row for row in model["comparisons"][0]["metrics"] if row["metric"] == "ICMP echo nonresponse (%)")
        self.assertIsNone(percentage["after"])
        self.assertIsNone(percentage["delta"])
        after["settings"]["count"] = 5
        after["result"]["summary"] = summarize_check("loss", after["result"], after["settings"])
        self.manager.save(after)
        with self.assertRaisesRegex(ValueError, "changed"):
            prepare_report(self.manager, {"ids": [before["id"], after["id"]], "fingerprint": model["fingerprint"]})

    def test_json_xml_import_reopen_and_report_do_not_run_network_work(self):
        from unittest.mock import patch
        record = self.save("dns")
        source = self.manager.export(record["id"])
        for xml in (False, True):
            with tempfile.TemporaryDirectory() as directory:
                other = RunManager(directory, "missing-iperf")
                try:
                    payload = xml_export(source) if xml else json.dumps(source).encode()
                    with patch("app.checks.CheckCapture.run", side_effect=AssertionError("Imports must never execute checks")):
                        result = import_runs(other, parse_document(payload), "fixture.xml" if xml else "fixture.json")
                        identity = result["imported"][0]
                        model = prepare_report(other, {"ids": [identity], "redact": False})
                    self.assertEqual(model["runs"][0]["result"]["summary"], record["result"]["summary"])
                    other.shutdown()
                    other = RunManager(directory, "missing-iperf")
                    self.assertEqual(other.read(identity)["result"], record["result"])
                finally:
                    other.shutdown()


if __name__ == "__main__":
    unittest.main()
