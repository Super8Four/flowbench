import copy
import base64
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import json
import re
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
import zlib

from docx import Document

from app.exports import xml_export
from app.flent import normalize as normalize_flent, RRUL_TCP, RRUL_RTT
from app.diagnostics import latency_summary
from app.imports import import_runs, parse_document
from app.projects import ProjectStore
from app.reports import prepare_report, render_docx, render_pdf, _chart
from app.runner import RunManager, settings_from


def decoded_pdf_streams(payload):
    result = []
    for stream in re.findall(rb"stream\r?\n(.*?)endstream", payload, re.S):
        try:
            decoded = base64.a85decode(stream.rstrip(b"\r\n"), adobe=True)
        except (ValueError, OverflowError):
            decoded = stream
        try:
            decoded = zlib.decompress(decoded)
        except zlib.error:
            pass
        result.append(decoded)
    return b"\n".join(result)


def report_fixture():
    baseline = [{"status": "reply", "rtt_ms": 2.1, "rtt_upper_bound_ms": None} for _ in range(3)]
    loaded = [{"status": "reply", "rtt_ms": 4.2, "rtt_upper_bound_ms": None} for _ in range(6)]
    return {"schema_version": 1, "id": "a" * 32, "created_at": "2026-01-01T10:00:00+00:00",
            "finished_at": "2026-01-01T10:00:10+00:00", "status": "completed",
            "settings": settings_from({"target": "192.0.2.10", "bind": "192.0.2.11", "notes": "private.example ticket",
                                       "label": "Private circuit", "tags": {"site": "private.example", "ticket": "Secret ticket"}}),
            "iperf_version": "iperf 3.22", "output_format": "json", "exit_code": 0,
            "command": ["iperf3", "-c", "192.0.2.10"], "error": None,
            "summary": {"sum_received": {"bits_per_second": 100e6}, "sum_sent": {"retransmits": 5}},
            "intervals": [{"sum": {"start": 0, "end": 1, "bits_per_second": 100e6}},
                          {"sum": {"start": 3, "end": 4, "bits_per_second": 120e6}}], "test_parameters": None,
            "annotations": [{"id": "b" * 32, "timestamp": "2026-01-01T10:00:03+00:00", "seconds": 3, "text": "Private note <tag>& caf\u00e9"}],
            "route": {"status": "captured", "target": "192.0.2.10", "raw_output": "private.example",
                      "hops": [{"hop": 1, "address": "192.0.2.1", "rtt_ms": [1.2, 1.3], "detail": "private.example", "responded": True}]},
            "diagnostics": {"schema_version": 1, "latency": {"enabled": True, "target": "192.0.2.10", "metric": "ICMP echo RTT",
                            "baseline": {"status": "captured", "summary": latency_summary(baseline), "samples": baseline},
                            "loaded": {"status": "captured", "summary": latency_summary(loaded), "samples": loaded},
                            "mean_increase_ms": 2.1, "error": "private.example timed out", "raw_stdout": "192.0.2.10"},
                            "endpoint_health": {"enabled": True, "local": {"provenance": "detected local", "delta": {"cpu_busy_percent": 15.2}}, "remote": {"status": "unavailable"}},
                            "path_context": {"mtu": {"status": "disabled", "tested_packet_sizes": [], "largest_confirmed_packet_bytes": None}}}}


def native_rrul_fixture():
    return {"version": 4, "metadata": {"NAME": "rrul", "HOST": "192.0.2.10", "LENGTH": 10},
            "x_values": [0, 1, 2, 3],
            "results": {name: [index + 1, None, index + 3, index + 4]
                        for index, name in enumerate(sorted(RRUL_TCP | RRUL_RTT))}}


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.manager = RunManager(self.temp.name, "missing-iperf")
        self.before = report_fixture()
        self.project = ProjectStore(self.manager).create({"name": "Private project", "description": "private.example"})
        self.before["settings"]["project_id"] = self.project["id"]
        self.manager.save(self.before)

    def tearDown(self):
        self.manager.shutdown()
        self.temp.cleanup()

    def test_default_redaction_overrides_raw_and_retains_actual_measurements(self):
        self.before["diagnostics"]["latency"]["private.example"] = "Secret ticket"
        self.manager.save(self.before)
        model = prepare_report(self.manager, {"ids": [self.before["id"]], "include_raw": True})
        serialized = json.dumps(model)
        for private in ("192.0.2.10", "192.0.2.11", "192.0.2.1", "private.example", "Private circuit", "Private project", "Private note", "Secret ticket"):
            self.assertNotIn(private, serialized)
        self.assertEqual(model["runs"][0]["metrics"]["upload_mbps"], 100)
        self.assertEqual(model["runs"][0]["diagnostics"]["latency"]["loaded"]["summary"]["mean_ms"], 4.2)
        self.assertEqual(model["runs"][0]["raw_evidence"], {"status": "removed by redaction"})
        self.assertEqual(model["runs"][0]["annotations"][0]["seconds"], 3)

    def test_private_context_requires_opt_out_and_raw_separate_opt_in(self):
        model = prepare_report(self.manager, {"ids": [self.before["id"]], "redact": False})
        self.assertEqual(model["runs"][0]["settings"]["target"], "192.0.2.10")
        self.assertEqual(model["runs"][0]["project"]["name"], "Private project")
        self.assertNotIn("raw_evidence", model["runs"][0])
        self.assertNotIn("raw_output", model["runs"][0]["route"])
        self.assertNotIn("raw_stdout", model["runs"][0]["diagnostics"]["latency"])
        raw = prepare_report(self.manager, {"ids": [self.before["id"]], "redact": False, "include_raw": True})
        self.assertIn("raw_evidence", raw["runs"][0])
        self.assertEqual(raw["runs"][0]["route"]["raw_output"], "private.example")

    def test_gates_invalid_options_and_incomplete_or_parent_records(self):
        for options in ({}, {"ids": []}, {"ids": ["invalid"]}, {"ids": [self.before["id"]] * 2},
                        {"ids": [self.before["id"]], "redact": "false"}, {"ids": [self.before["id"]], "extra": 1}):
            with self.assertRaises(ValueError):
                prepare_report(self.manager, options)
        for change in ({"status": "running"}, {"status": "failed"}, {"status": "cancelled"},
                       {"kind": "matrix-job"}, {"kind": "protocol-job"}, {"summary": {}}, {"finished_at": None}):
            self.manager.save(dict(self.before, **change))
            with self.assertRaises(ValueError):
                prepare_report(self.manager, {"ids": [self.before["id"]]})
        with self.assertRaises(KeyError):
            prepare_report(self.manager, {"ids": ["f" * 32]})

    def test_comparison_warns_and_zero_baseline_percent_remains_undefined(self):
        self.before["summary"]["sum_sent"]["retransmits"] = 0
        self.manager.save(self.before)
        after = copy.deepcopy(self.before)
        after["id"] = "c" * 32
        after["settings"]["streams"] = 2
        after["summary"]["sum_received"]["bits_per_second"] = 200e6
        after["summary"]["sum_sent"]["retransmits"] = 3
        self.manager.save(after)
        model = prepare_report(self.manager, {"ids": [self.before["id"], after["id"]]})
        report = model["comparisons"][0]
        self.assertIn("streams", report["settings_differences"])
        self.assertIn("do not establish", report["caveat"])
        self.assertIsNone(next(item for item in report["metrics"] if item["metric"] == "upload_retransmits")["percent_change"])

    def test_word_comparison_heading_explicitly_stays_with_first_metric(self):
        after = copy.deepcopy(self.before)
        after["id"] = "c" * 32
        after["settings"]["streams"] = 2
        self.manager.save(after)
        model = prepare_report(self.manager, {"ids": [self.before["id"], after["id"]]})
        document = Document(BytesIO(render_docx(model)))
        heading = next(p for p in document.paragraphs if p.text == "Baseline compared with Test 2")
        # Both inherited and direct keepNext were insufficient in actual AMD64
        # office conversion. Require an explicit page start, plus the full chain
        # through both caveats and the table header to its first measured row.
        self.assertIs(heading.paragraph_format.page_break_before, True)
        self.assertIs(heading.paragraph_format.keep_with_next, True)
        self.assertIs(heading.paragraph_format.keep_together, True)
        following = heading._p.getnext()
        for _ in range(2):
            self.assertTrue(following.xpath("./w:pPr/w:keepNext"))
            following = following.getnext()
        self.assertTrue(following.tag.endswith("}tbl"))
        self.assertTrue(following.xpath("./w:tr[1]/w:tc/w:p/w:pPr/w:keepNext"))

    def test_rendered_formats_escape_text_and_are_repeatable_without_archive_writes(self):
        original = self.manager.export(self.before["id"])
        model = prepare_report(self.manager, {"ids": [self.before["id"]], "redact": False})
        for _ in range(2):
            pdf = render_pdf(model)
            self.assertTrue(pdf.startswith(b"%PDF-"))
            self.assertIn(b"/FontFile2", pdf)
            self.assertIn(b"/EmbeddedFiles", pdf)
            self.assertIn(b"Copyright (c) 2003 by Bitstream", pdf)
            word = render_docx(model)
            with ZipFile(BytesIO(word)) as archive:
                self.assertIn("word/document.xml", archive.namelist())
                self.assertIn(b"Private note &lt;tag&gt;&amp;", archive.read("word/document.xml"))
                self.assertFalse(any("vba" in name.lower() for name in archive.namelist()))
                relationships = b"".join(archive.read(name) for name in archive.namelist() if name.endswith(".rels"))
                self.assertNotIn(b'TargetMode="External"', relationships)
            document = Document(BytesIO(word))
            self.assertEqual(document.core_properties.author, "FlowBench")
            self.assertTrue(any("Mean increase ms" in cell.text for table in document.tables for row in table.rows for cell in row.cells))
        self.assertEqual(original, self.manager.export(self.before["id"]))

    def test_chart_does_not_bridge_missing_intervals(self):
        from PIL import Image
        from PIL import ImageColor
        model = prepare_report(self.manager, {"ids": [self.before["id"]]})
        image = Image.open(BytesIO(_chart(model["runs"])))
        # Seconds 2 lies in the uncaptured span between [0,1] and [3,4].
        x = int(130 + 2 / 4 * 1305)
        for y in range(80, 460):
            self.assertNotEqual(image.getpixel((x, y)), ImageColor.getrgb(model["runs"][0]["color"]))

    def test_sqlite_restart_and_json_xml_roundtrip_preserve_report_evidence(self):
        self.before["settings"]["color"] = "#ABCDEF"
        self.manager.save(self.before)
        self.manager.shutdown()
        self.manager = RunManager(self.temp.name, "missing-iperf")
        original = self.manager.export(self.before["id"])
        for xml in (False, True):
            with tempfile.TemporaryDirectory() as other:
                restored = RunManager(other, "missing-iperf")
                try:
                    document = parse_document(xml_export(original)) if xml else parse_document(json.dumps(original).encode())
                    imported = import_runs(restored, document, "fixture.xml" if xml else "fixture.json")
                    model = prepare_report(restored, {"ids": imported["imported"], "redact": False, "include_raw": True})
                    row = model["runs"][0]
                    self.assertEqual(row["diagnostics"], original["diagnostics"])
                    self.assertEqual(row["route"]["hops"], original["route"]["hops"])
                    self.assertEqual(row["annotations"], original["annotations"])
                    self.assertTrue(row["provenance"]["imported"])
                    self.assertEqual(row["color"], "#abcdef")
                finally:
                    restored.shutdown()

    def test_external_series_and_metric_names_are_redacted(self):
        native = native_rrul_fixture()
        native["results"]["TCP upload private.example"] = [10, None, 30, 40]
        native["results"]["private-unit-series.example"] = [1, None, 2, 3]
        native["metadata"]["SERIES_META"] = {
            "TCP upload private.example": {"UNITS": "Mbits/s"},
            "private-unit-series.example": {"UNITS": "private-unit.example"}}
        source = dict(self.before, kind="flent-run", summary={}, intervals=[],
                      flent=normalize_flent(native, "private.example version"))
        self.manager.save(source)
        model = prepare_report(self.manager, {"ids": [source["id"]], "include_raw": True})
        self.assertNotIn("private.example", json.dumps(model))
        self.assertNotIn("private-unit", json.dumps(model))
        self.assertEqual(model["runs"][0]["flent_series"][-2]["points"][1]["value"], None)
        self.assertEqual(model["runs"][0]["flent_summaries"][-2]["missing"], 1)
        self.assertIsNone(model["runs"][0]["flent_series"][-1]["unit"])
        for output in (decoded_pdf_streams(render_pdf(model)), render_docx(model)):
            if output.startswith(b"PK"):
                with ZipFile(BytesIO(output)) as package:
                    output = b"\n".join(package.read(name) for name in package.namelist() if name.endswith(".xml"))
            self.assertNotIn(b"private.example", output)
            self.assertNotIn(b"private-unit", output)

    def test_historical_completed_flent_rows_require_complete_native_rrul_without_mutation(self):
        from app import create_app
        with tempfile.TemporaryDirectory() as archive:
            app = create_app({"TESTING": True, "RUNS_DIRECTORY": archive, "IPERF_PATH": "missing-iperf"})
            manager = app.extensions["iperf_runs"]
            try:
                native = native_rrul_fixture()
                complete = dict(self.before, kind="flent-run", summary={}, intervals=[],
                                flent=normalize_flent(native, "Synthetic fixture"))
                manager.save(complete)
                client = app.test_client()
                self.assertEqual(client.post("/api/reports/preview", json={"ids": [complete["id"]]}).status_code, 200)
                damaged = copy.deepcopy(native)
                for name in RRUL_TCP:
                    damaged["results"][name] = [None] * len(damaged["x_values"])
                partial = normalize_flent(damaged, "Synthetic fixture")
                for evidence in (dict(partial, status="captured"), partial,
                                 dict(complete["flent"], data=None),
                                 dict(complete["flent"], summaries=[])):
                    with self.subTest(evidence_status=evidence["status"], native_source=evidence.get("data") is not None):
                        # Direct storage models an older SQLite row; no import can
                        # manufacture today's impossible completed evidence.
                        legacy = dict(complete, flent=evidence)
                        manager.save(legacy)
                        original = manager.export(legacy["id"])
                        for endpoint in ("/api/reports/preview", "/api/reports/download?format=pdf", "/api/reports/download?format=docx"):
                            response = client.post(endpoint, json={"ids": [legacy["id"]]})
                            self.assertEqual(response.status_code, 400)
                            self.assertIn("complete native Flent RRUL", response.get_json()["error"])
                        self.assertEqual(original, manager.export(legacy["id"]))
            finally:
                manager.shutdown()

    def test_preview_fingerprint_accepts_same_snapshot_and_rejects_changes(self):
        options = {"ids": [self.before["id"]]}
        first = prepare_report(self.manager, options)
        second = prepare_report(self.manager, dict(options, fingerprint=first["fingerprint"]))
        self.assertEqual(first["fingerprint"], second["fingerprint"])
        changed = copy.deepcopy(self.before)
        changed["summary"]["sum_received"]["bits_per_second"] *= 2
        self.manager.save(changed)
        with self.assertRaisesRegex(ValueError, "Saved evidence changed"):
            prepare_report(self.manager, dict(options, fingerprint=first["fingerprint"]))

    def test_concurrent_report_requests_keep_archive_and_models_unchanged(self):
        before = self.manager.export(self.before["id"])
        def generate(index):
            model = prepare_report(self.manager, {"ids": [self.before["id"]]})
            pdf, word = render_pdf(model), render_docx(model)
            self.assertTrue(pdf.startswith(b"%PDF-"))
            with ZipFile(BytesIO(word)) as package:
                self.assertIsNone(package.testzip())
            return model["fingerprint"]
        with ThreadPoolExecutor(max_workers=4) as pool:
            fingerprints = list(pool.map(generate, range(4)))
        self.assertEqual(len(set(fingerprints)), 1)
        self.assertEqual(self.manager.export(self.before["id"]), before)

    def test_ten_runs_preserve_colors_canonical_rates_and_shared_window(self):
        ids = []
        for index in range(10):
            source = copy.deepcopy(self.before)
            source["id"] = format(index + 1, "032x")
            source["settings"].update(duration=30 if index == 0 else 10, color="#123456" if index == 0 else "#b52b3d")
            source["finished_at"] = "2026-01-01T10:00:30+00:00" if index == 0 else "2026-01-01T10:00:10+00:00"
            source["summary"]["sum_received"]["bits_per_second"] = (0 if index == 0 else 3000) * 1e6
            source["intervals"] = [{"sum": {"start": seconds, "end": seconds + 1, "bits_per_second": (0 if index == 0 else 3000) * 1e6}} for seconds in range(source["settings"]["duration"])]
            self.manager.save(source)
            ids.append(source["id"])
        model = prepare_report(self.manager, {"ids": ids})
        self.assertEqual(len(model["runs"]), 10)
        self.assertEqual(len(model["comparisons"]), 9)
        self.assertEqual(model["runs"][0]["color"], "#123456")
        self.assertEqual(model["runs"][1]["metrics"]["upload_mbps"], 3000)
        self.assertEqual(model["runs"][1]["series"][-1]["seconds"], 10)
        self.assertEqual(model["runs"][0]["series"][-1]["seconds"], 30)
        self.assertEqual(model["common_window"]["seconds"], 10)
        self.assertEqual(model["display"]["rate_axis"]["unit"], "Gbps")
        word = Document(BytesIO(render_docx(model)))
        text = " ".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
        self.assertIn("3 Gbps", text)
        self.assertIn("0 bps", text)
        self.assertIn("#123456", text)

    def test_resource_limits_depend_on_data_not_selection_count(self):
        with patch("app.reports.CHART_POINTS", 1):
            with self.assertRaisesRegex(ValueError, "measured-point resource budget"):
                prepare_report(self.manager, {"ids": [self.before["id"]]})
        with patch("app.reports.MODEL_BYTES", 1):
            with self.assertRaisesRegex(ValueError, "normalized-data resource budget"):
                prepare_report(self.manager, {"ids": [self.before["id"]]})

    def test_json_xml_import_cannot_leak_custom_directions_or_streams_in_redacted_reports(self):
        from app import create_app
        direction_secret = "private-direction.example"
        stream_secret = "private-stream.example"
        for use_xml in (False, True):
            with self.subTest(xml=use_xml), tempfile.TemporaryDirectory() as archive:
                app = create_app({"TESTING": True, "RUNS_DIRECTORY": archive, "IPERF_PATH": "missing-iperf"})
                manager = app.extensions["iperf_runs"]
                client = app.test_client()
                try:
                    before = report_fixture()
                    before["intervals"] = [{"start": 0, "end": 1, "bits_per_second": 100e6,
                                             "direction": direction_secret, "stream": stream_secret}]
                    after = copy.deepcopy(before)
                    after["id"] = "c" * 32
                    after["intervals"][0]["direction"] = "upload"
                    after["summary"]["sum_received"]["bits_per_second"] = 200e6
                    bundle = {"schema_version": 1, "before": before, "after": after}
                    content = xml_export(bundle).decode() if use_xml else json.dumps(bundle)
                    imported = client.post("/api/runs/import", json={"filename": "fixture.xml" if use_xml else "fixture.json", "content": content})
                    self.assertEqual(imported.status_code, 200, imported.get_json())
                    ids = imported.get_json()["imported"]
                    originals = [manager.export(run_id) for run_id in ids]
                    preview = client.post("/api/reports/preview", json={"ids": ids, "include_raw": True})
                    self.assertEqual(preview.status_code, 200, preview.get_json())
                    model = preview.get_json()
                    for secret in (direction_secret, stream_secret):
                        self.assertNotIn(secret, json.dumps(model))
                    self.assertEqual(model["runs"][0]["series"][0]["mbps"], 100)
                    self.assertEqual(model["runs"][1]["series"][0]["direction"], "upload")
                    self.assertEqual(model["common_window"]["runs"][0]["directions"][0]["mean_mbps"], 100)
                    private = client.post("/api/reports/preview", json={"ids": ids, "redact": False}).get_json()
                    self.assertIn(direction_secret, json.dumps(private))
                    self.assertIn(stream_secret, json.dumps(private))
                    # The positive control proves decoded PDF streams catch this text.
                    self.assertIn(direction_secret.encode(), decoded_pdf_streams(render_pdf(private)))
                    with ZipFile(BytesIO(render_docx(private))) as package:
                        self.assertIn(direction_secret.encode(), package.read("word/document.xml"))
                    for format_name in ("pdf", "docx"):
                        downloaded = client.post("/api/reports/download?format=" + format_name,
                                                 json={"ids": ids, "include_raw": True, "fingerprint": model["fingerprint"]})
                        self.assertEqual(downloaded.status_code, 200)
                        if format_name == "pdf":
                            readable = decoded_pdf_streams(downloaded.data)
                        else:
                            with ZipFile(BytesIO(downloaded.data)) as package:
                                readable = b"\n".join(package.read(name) for name in package.namelist() if name.endswith(".xml"))
                        for secret in (direction_secret, stream_secret):
                            self.assertNotIn(secret.encode(), readable)
                    self.assertEqual(originals, [manager.export(run_id) for run_id in ids])
                    self.assertEqual(manager.export(ids[0])["intervals"][0]["direction"], direction_secret)
                finally:
                    manager.shutdown()
if __name__ == "__main__":
    unittest.main()
